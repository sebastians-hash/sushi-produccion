from flask import Flask, request, jsonify, send_file, render_template, session, redirect, url_for
import json, os, tempfile, functools
import psycopg2
import psycopg2.extras
from datetime import datetime
from gen_pdf import build_pdf
from gen_xlsx import build_insumos_xlsx
from parse_receta import parse_receta_pdf, match_insumo, formato_to_unidad
from authlib.integrations.flask_client import OAuth
from werkzeug.middleware.proxy_fix import ProxyFix

app = Flask(__name__)
app.secret_key = os.environ.get('FLASK_SECRET_KEY', 'CAMBIAR-ESTA-CLAVE-EN-PRODUCCION')
# Railway (y la mayoria de plataformas cloud) terminan el HTTPS en su proxy y
# reenvian a la app como HTTP interno. Sin esto, url_for(..., _external=True)
# generaria URLs http:// en vez de https://, rompiendo el callback de Google OAuth.
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

# DATABASE_URL la provee Railway automaticamente al agregar un servicio de PostgreSQL
# y vincularlo a esta app. Es una base de datos administrada: los datos NO se pierden
# con cada deploy, a diferencia de un archivo SQLite en el filesystem del contenedor.
DATABASE_URL = os.environ.get('DATABASE_URL')

# ── Postgres IntegrityError, expuesto con el mismo nombre que usa el resto del codigo ──
IntegrityError = psycopg2.IntegrityError

# ── DB ────────────────────────────────────────────────────────────────────
class _PGCursorWrapper:
    """Envuelve un cursor de psycopg2 para que .execute() acepte '?' como
    placeholder (estilo sqlite3) y devuelva filas con acceso tipo diccionario,
    manteniendo compatible el resto del codigo sin reescribir cada consulta."""
    def __init__(self, cursor):
        self._cursor = cursor
    def execute(self, sql, params=()):
        sql_pg = sql.replace('?', '%s')
        self._cursor.execute(sql_pg, params)
        return self
    def executescript(self, script):
        self._cursor.execute(script)
        return self
    def fetchone(self):
        return self._cursor.fetchone()
    def fetchall(self):
        return self._cursor.fetchall()
    @property
    def lastrowid(self):
        return None  # no usado; los INSERT que necesitan el id usan RETURNING

class _PGConnWrapper:
    def __init__(self, conn):
        self._conn = conn
    def execute(self, sql, params=()):
        cur = self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        sql_pg = sql.replace('?', '%s')
        cur.execute(sql_pg, params)
        return _PGCursorWrapper(cur)
    def executescript(self, script):
        cur = self._conn.cursor()
        cur.execute(script)
        return _PGCursorWrapper(cur)
    def cursor(self):
        cur = self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        return _PGCursorWrapper(cur)
    def commit(self):
        self._conn.commit()
    def rollback(self):
        self._conn.rollback()
    def close(self):
        self._conn.close()

def get_db():
    conn = psycopg2.connect(DATABASE_URL)
    return _PGConnWrapper(conn)

# ── Google OAuth setup ──────────────────────────────────────────────────────
oauth = OAuth(app)
google = oauth.register(
    name='google',
    client_id=os.environ.get('GOOGLE_CLIENT_ID'),
    client_secret=os.environ.get('GOOGLE_CLIENT_SECRET'),
    server_metadata_url='https://accounts.google.com/.well-known/openid-configuration',
    client_kwargs={'scope': 'openid email profile'}
)

def init_db():
    conn = get_db()
    c = conn.cursor()
    c.executescript('''
        CREATE TABLE IF NOT EXISTS rolls (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            insumos TEXT NOT NULL,
            piezas_por_rollo INTEGER NOT NULL DEFAULT 14,
            rollo_blanco_grupo TEXT,
            cuenta_productividad BOOLEAN NOT NULL DEFAULT TRUE,
            toppings TEXT NOT NULL DEFAULT '[]'
        );
        CREATE TABLE IF NOT EXISTS rollo_blanco_grupos (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL
        );
        CREATE TABLE IF NOT EXISTS combos (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            family TEXT NOT NULL,
            rolls TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS semielaborados (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            insumo_key TEXT NOT NULL,
            unit TEXT NOT NULL,
            rolls TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS sushimanes (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            productivity INTEGER NOT NULL DEFAULT 10,
            active INTEGER NOT NULL DEFAULT 1,
            dias_franco TEXT NOT NULL DEFAULT '[]',
            horario_ingreso TEXT,
            turno TEXT,
            posicion_id INTEGER
        );
        CREATE TABLE IF NOT EXISTS posiciones (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            orden INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS insumos (
            id SERIAL PRIMARY KEY,
            key TEXT UNIQUE NOT NULL,
            label TEXT NOT NULL,
            unidad_receta TEXT NOT NULL DEFAULT 'g',
            unidad_resumen TEXT NOT NULL DEFAULT 'kg',
            factor_conversion REAL NOT NULL DEFAULT 0.001,
            precio_unidad REAL,
            categoria TEXT,
            es_80_20 INTEGER NOT NULL DEFAULT 0,
            comentario TEXT,
            marca_producto TEXT,
            marca_tipo TEXT,
            zona_almacenamiento TEXT,
            proveedor_principal_id INTEGER,
            proveedor_alt1_id INTEGER,
            proveedor_alt2_id INTEGER,
            eficiencia REAL NOT NULL DEFAULT 100
        );
        CREATE TABLE IF NOT EXISTS categorias_insumos (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL
        );
        CREATE TABLE IF NOT EXISTS zonas_almacenamiento (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL
        );
        CREATE TABLE IF NOT EXISTS unidades (
            id SERIAL PRIMARY KEY,
            nombre TEXT NOT NULL,
            simbolo TEXT UNIQUE NOT NULL,
            orden INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS proveedores (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            contactos TEXT NOT NULL DEFAULT '[]'
        );
        CREATE TABLE IF NOT EXISTS marcas (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            orden INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS usuarios (
            id SERIAL PRIMARY KEY,
            email TEXT UNIQUE NOT NULL,
            nombre TEXT,
            role TEXT NOT NULL DEFAULT 'user',
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT
        );
        CREATE TABLE IF NOT EXISTS familias (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL
        );
        CREATE TABLE IF NOT EXISTS otros_productos (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            tipo TEXT NOT NULL DEFAULT 'porcion',
            familia TEXT,
            insumos TEXT NOT NULL DEFAULT '{}',
            marcas TEXT NOT NULL DEFAULT '[]',
            rolls TEXT NOT NULL DEFAULT '{}'
        );
        CREATE TABLE IF NOT EXISTS equivalencias (
            id SERIAL PRIMARY KEY,
            tipo TEXT NOT NULL,
            nombre_canonico TEXT NOT NULL,
            nombre_alias TEXT NOT NULL,
            marca TEXT,
            factor REAL NOT NULL DEFAULT 1.0,
            UNIQUE(tipo, nombre_alias, marca)
        );
        CREATE TABLE IF NOT EXISTS sugerencias (
            id SERIAL PRIMARY KEY,
            tipo TEXT NOT NULL,
            titulo TEXT NOT NULL,
            descripcion TEXT NOT NULL,
            estado TEXT NOT NULL DEFAULT 'nueva',
            respuesta_admin TEXT,
            creado_por TEXT NOT NULL,
            creado_por_nombre TEXT,
            referencia_tipo TEXT,
            referencia_id INTEGER,
            referencia_nombre TEXT,
            visto BOOLEAN NOT NULL DEFAULT TRUE,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS novedades (
            id SERIAL PRIMARY KEY,
            titulo TEXT NOT NULL,
            desarrollo TEXT NOT NULL,
            creado_por TEXT NOT NULL,
            creado_por_nombre TEXT,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS novedades_leidas (
            id SERIAL PRIMARY KEY,
            novedad_id INTEGER NOT NULL,
            user_email TEXT NOT NULL,
            leida_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(novedad_id, user_email)
        );
        CREATE TABLE IF NOT EXISTS locales (
            id SERIAL PRIMARY KEY,
            name TEXT UNIQUE NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            marcas TEXT NOT NULL DEFAULT '[]'
        );
        CREATE TABLE IF NOT EXISTS usuario_locales (
            id SERIAL PRIMARY KEY,
            usuario_id INTEGER NOT NULL,
            local_id INTEGER NOT NULL,
            UNIQUE(usuario_id, local_id)
        );
        CREATE TABLE IF NOT EXISTS planillas (
            id SERIAL PRIMARY KEY,
            local_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            estado TEXT NOT NULL DEFAULT 'borrador',
            data TEXT NOT NULL DEFAULT '{}',
            created_by TEXT,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
    ''')
    conn.commit()

    # Bootstrap: si no hay locales todavia, crear uno por defecto y migrar
    # a el todo lo que ya existia (para no romper la operacion actual de un
    # solo local mientras se suman los demas)
    n_locales = c.execute('SELECT COUNT(*) AS cnt FROM locales').fetchone()['cnt']
    default_local_id = None
    if n_locales == 0:
        c.execute("INSERT INTO locales (name, active) VALUES ('Local Principal', 1)")
        conn.commit()
        default_local_id = c.execute("SELECT id FROM locales WHERE name='Local Principal'").fetchone()['id']

    # Bootstrap: si no hay usuarios todavía, crear el admin inicial
    # a partir de la variable de entorno ADMIN_EMAIL
    admin_email = os.environ.get('ADMIN_EMAIL')
    n_users = c.execute('SELECT COUNT(*) AS cnt FROM usuarios').fetchone()['cnt']
    if n_users == 0 and admin_email:
        c.execute('INSERT INTO usuarios (email, nombre, role, active, created_at) VALUES (?,?,?,1,?) ON CONFLICT (email) DO NOTHING',
                   (admin_email.strip().lower(), 'Administrador', 'admin', datetime.now().isoformat()))
        conn.commit()

    # Sembrar familias a partir de las que ya usan los combos existentes
    # (para no romper nada) mas categorias sugeridas para los productos nuevos
    n_familias = c.execute('SELECT COUNT(*) AS cnt FROM familias').fetchone()['cnt']
    if n_familias == 0:
        existing_families = [r['family'] for r in c.execute(
            "SELECT DISTINCT family FROM combos WHERE family IS NOT NULL AND family != ''").fetchall()]
        suggested = ['Porciones', 'Ensaladas', 'Entradas', 'Platos Calientes', 'Otros']
        all_families = list(dict.fromkeys(existing_families + suggested))  # dedup preservando orden
        for fam in all_families:
            c.execute('INSERT INTO familias (name) VALUES (?) ON CONFLICT (name) DO NOTHING', (fam,))
        conn.commit()

    # Migration: add recipe columns to semielaborados if they don't exist yet
    def get_columns(table):
        rows = c.execute(
            "SELECT column_name FROM information_schema.columns WHERE table_name = ?",
            (table,)).fetchall()
        return [r['column_name'] for r in rows]

    existing_cols = get_columns('semielaborados')
    if 'receta' not in existing_cols:
        c.execute("ALTER TABLE semielaborados ADD COLUMN receta TEXT NOT NULL DEFAULT '[]'")
    if 'rendimiento_cantidad' not in existing_cols:
        c.execute("ALTER TABLE semielaborados ADD COLUMN rendimiento_cantidad REAL NOT NULL DEFAULT 0")
    if 'rendimiento_unidad' not in existing_cols:
        c.execute("ALTER TABLE semielaborados ADD COLUMN rendimiento_unidad TEXT NOT NULL DEFAULT 'g'")
    if 'marcas' not in existing_cols:
        c.execute("ALTER TABLE semielaborados ADD COLUMN marcas TEXT NOT NULL DEFAULT '[]'")
    if 'tiempo_elaboracion_min' not in existing_cols:
        c.execute("ALTER TABLE semielaborados ADD COLUMN tiempo_elaboracion_min INTEGER")
    if 'vida_util_dias' not in existing_cols:
        c.execute("ALTER TABLE semielaborados ADD COLUMN vida_util_dias INTEGER")

    otro_cols = get_columns('otros_productos')
    if 'rolls' not in otro_cols:
        c.execute("ALTER TABLE otros_productos ADD COLUMN rolls TEXT NOT NULL DEFAULT '{}'")
    if 'procedimiento' not in otro_cols:
        c.execute("ALTER TABLE otros_productos ADD COLUMN procedimiento TEXT NOT NULL DEFAULT '[]'")

    equiv_cols = get_columns('equivalencias')
    if 'factor' not in equiv_cols:
        c.execute("ALTER TABLE equivalencias ADD COLUMN factor REAL NOT NULL DEFAULT 1.0")

    # Migration: add 'marcas' column to combos and rolls
    combo_cols = get_columns('combos')
    if 'marcas' not in combo_cols:
        c.execute("ALTER TABLE combos ADD COLUMN marcas TEXT NOT NULL DEFAULT '[]'")

    roll_cols = get_columns('rolls')
    if 'marcas' not in roll_cols:
        c.execute("ALTER TABLE rolls ADD COLUMN marcas TEXT NOT NULL DEFAULT '[]'")
    if 'piezas_por_rollo' not in roll_cols:
        c.execute("ALTER TABLE rolls ADD COLUMN piezas_por_rollo INTEGER NOT NULL DEFAULT 14")
    if 'rollo_blanco_grupo' not in roll_cols:
        c.execute("ALTER TABLE rolls ADD COLUMN rollo_blanco_grupo TEXT")
    if 'cuenta_productividad' not in roll_cols:
        c.execute("ALTER TABLE rolls ADD COLUMN cuenta_productividad BOOLEAN NOT NULL DEFAULT TRUE")
    if 'toppings' not in roll_cols:
        c.execute("ALTER TABLE rolls ADD COLUMN toppings TEXT NOT NULL DEFAULT '[]'")

    local_cols = get_columns('locales')
    if 'marcas' not in local_cols:
        c.execute("ALTER TABLE locales ADD COLUMN marcas TEXT NOT NULL DEFAULT '[]'")

    marca_cols = get_columns('marcas')
    if 'orden' not in marca_cols:
        c.execute("ALTER TABLE marcas ADD COLUMN orden INTEGER NOT NULL DEFAULT 0")
        c.execute("UPDATE marcas SET orden = id")  # respeta el orden en que se fueron creando, como punto de partida

    # La restricción vieja no dejaba repetir un mismo nombre de venta entre marcas
    # distintas, aunque fuera para el mismo producto — la reemplazamos por una que
    # solo exige que sea único DENTRO de cada marca.
    c.execute("ALTER TABLE equivalencias DROP CONSTRAINT IF EXISTS equivalencias_tipo_nombre_alias_key")
    c.execute("""DO $$ BEGIN
        IF NOT EXISTS (
            SELECT 1 FROM pg_constraint WHERE conname = 'equivalencias_tipo_nombre_alias_marca_key'
        ) THEN
            ALTER TABLE equivalencias ADD CONSTRAINT equivalencias_tipo_nombre_alias_marca_key UNIQUE (tipo, nombre_alias, marca);
        END IF;
    END $$;""")

    # Sembrar las unidades que ya se usaban (antes fijas en el código) + las que se pidieron de arranque
    ya_hay_unidades = c.execute('SELECT COUNT(*) AS n FROM unidades').fetchone()['n']
    if not ya_hay_unidades:
        unidades_iniciales = [
            ('Gramos', 'g'), ('Mililitros', 'ml'), ('Unidades', 'u'), ('Hojas', 'hojas'),
            ('Kilogramos', 'kg'), ('Litros', 'L'),
            ('Paquetes', 'paq'), ('Manga', 'manga'),
        ]
        for idx, (nombre, simbolo) in enumerate(unidades_iniciales):
            c.execute('INSERT INTO unidades (nombre, simbolo, orden) VALUES (?,?,?) ON CONFLICT (simbolo) DO NOTHING', (nombre, simbolo, idx))

    sug_cols = get_columns('sugerencias')
    if 'referencia_tipo' not in sug_cols:
        c.execute("ALTER TABLE sugerencias ADD COLUMN referencia_tipo TEXT")
    if 'referencia_id' not in sug_cols:
        c.execute("ALTER TABLE sugerencias ADD COLUMN referencia_id INTEGER")
    if 'referencia_nombre' not in sug_cols:
        c.execute("ALTER TABLE sugerencias ADD COLUMN referencia_nombre TEXT")
    if 'visto' not in sug_cols:
        c.execute("ALTER TABLE sugerencias ADD COLUMN visto BOOLEAN NOT NULL DEFAULT TRUE")
    conn.commit()

    sm_cols = get_columns('sushimanes')
    if 'dias_franco' not in sm_cols:
        c.execute("ALTER TABLE sushimanes ADD COLUMN dias_franco TEXT NOT NULL DEFAULT '[]'")
    if 'horario_ingreso' not in sm_cols:
        c.execute("ALTER TABLE sushimanes ADD COLUMN horario_ingreso TEXT")
    if 'local_id' not in sm_cols:
        c.execute("ALTER TABLE sushimanes ADD COLUMN local_id INTEGER")
        # El nombre de un sushiman ya no tiene que ser unico en toda la empresa,
        # solo dentro de su propio local. Buscamos y sacamos la restriccion vieja
        # (unique sobre "name" sola) y ponemos una nueva sobre (name, local_id).
        old_constraints = c.execute("""
            SELECT tc.constraint_name FROM information_schema.table_constraints tc
            JOIN information_schema.constraint_column_usage ccu ON tc.constraint_name = ccu.constraint_name
            WHERE tc.table_name='sushimanes' AND tc.constraint_type='UNIQUE' AND ccu.column_name='name'
        """).fetchall()
        for row in old_constraints:
            c.execute(f'ALTER TABLE sushimanes DROP CONSTRAINT "{row["constraint_name"]}"')
        c.execute("ALTER TABLE sushimanes ADD CONSTRAINT sushimanes_name_local_unique UNIQUE (name, local_id)")
        # Asignar todos los sushimanes existentes al local por defecto (recien creado o el primero que exista)
        first_local = c.execute('SELECT id FROM locales ORDER BY id LIMIT 1').fetchone()
        if first_local:
            c.execute('UPDATE sushimanes SET local_id=? WHERE local_id IS NULL', (first_local['id'],))

    sm_cols = get_columns('sushimanes')
    if 'turno' not in sm_cols:
        c.execute("ALTER TABLE sushimanes ADD COLUMN turno TEXT")
    if 'posicion_id' not in sm_cols:
        c.execute("ALTER TABLE sushimanes ADD COLUMN posicion_id INTEGER")

    # Migration: nuevos atributos de insumos (categoria, 80/20, comentario, marca, proveedores, zona)
    insumo_cols = get_columns('insumos')
    insumo_new_cols = {
        'categoria': "ALTER TABLE insumos ADD COLUMN categoria TEXT",
        'es_80_20': "ALTER TABLE insumos ADD COLUMN es_80_20 INTEGER NOT NULL DEFAULT 0",
        'comentario': "ALTER TABLE insumos ADD COLUMN comentario TEXT",
        'marca_producto': "ALTER TABLE insumos ADD COLUMN marca_producto TEXT",
        'marca_tipo': "ALTER TABLE insumos ADD COLUMN marca_tipo TEXT",
        'zona_almacenamiento': "ALTER TABLE insumos ADD COLUMN zona_almacenamiento TEXT",
        'proveedor_principal_id': "ALTER TABLE insumos ADD COLUMN proveedor_principal_id INTEGER",
        'proveedor_alt1_id': "ALTER TABLE insumos ADD COLUMN proveedor_alt1_id INTEGER",
        'proveedor_alt2_id': "ALTER TABLE insumos ADD COLUMN proveedor_alt2_id INTEGER",
        'eficiencia': "ALTER TABLE insumos ADD COLUMN eficiencia REAL NOT NULL DEFAULT 100",
    }
    for col, stmt in insumo_new_cols.items():
        if col not in insumo_cols:
            c.execute(stmt)

    conn.commit()

    # Seed if empty
    if c.execute('SELECT COUNT(*) AS cnt FROM rolls').fetchone()['cnt'] == 0:
        seed_path = os.path.join(os.path.dirname(__file__), 'data', 'seed.json')
        if os.path.exists(seed_path):
            with open(seed_path) as f:
                seed = json.load(f)
            for roll in seed.get('rolls', []):
                c.execute('INSERT INTO rolls (name, insumos) VALUES (?,?) ON CONFLICT (name) DO NOTHING',
                          (roll['name'], json.dumps(roll['insumos'])))
            for combo in seed.get('combos', []):
                c.execute('INSERT INTO combos (name, family, rolls) VALUES (?,?,?) ON CONFLICT (name) DO NOTHING',
                          (combo['name'], combo['family'], json.dumps(combo['rolls'])))
            for semi in seed.get('semielaborados', []):
                c.execute('INSERT INTO semielaborados (name, insumo_key, unit, rolls) VALUES (?,?,?,?) ON CONFLICT (name) DO NOTHING',
                          (semi['name'], semi['insumo_key'], semi['unit'], json.dumps(semi['rolls'])))
            seed_local = c.execute('SELECT id FROM locales ORDER BY id LIMIT 1').fetchone()
            seed_local_id = seed_local['id'] if seed_local else None
            for sm in seed.get('sushimanes', []):
                c.execute('INSERT INTO sushimanes (name, productivity, local_id) VALUES (?,?,?) ON CONFLICT (name, local_id) DO NOTHING',
                          (sm['name'], sm['productivity'], seed_local_id))
            for ins in seed.get('insumos', []):
                c.execute('''INSERT INTO insumos
                    (key,label,unidad_receta,unidad_resumen,factor_conversion,precio_unidad)
                    VALUES (?,?,?,?,?,?) ON CONFLICT (key) DO NOTHING''',
                    (ins['key'], ins['label'], ins['unidad_receta'],
                     ins['unidad_resumen'], ins['factor_conversion'], ins.get('precio_unidad')))
            conn.commit()
    conn.close()

# ── Auth helpers ─────────────────────────────────────────────────────────
def login_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'user_email' not in session:
            if request.path.startswith('/api/'):
                return jsonify({'error': 'No autenticado'}), 401
            return redirect('/login')
        return f(*args, **kwargs)
    return decorated

def admin_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'user_email' not in session:
            if request.path.startswith('/api/'):
                return jsonify({'error': 'No autenticado'}), 401
            return redirect('/login')
        if session.get('user_role') != 'admin':
            return jsonify({'error': 'Requiere permisos de administrador'}), 403
        return f(*args, **kwargs)
    return decorated

import traceback

@app.errorhandler(Exception)
def handle_error(e):
    if request.path.startswith('/api/'):
        traceback.print_exc()  # queda en los logs de Railway para diagnostico
        return jsonify({'error': f'Error interno: {e}'}), 500
    raise e

# ── Auth routes ──────────────────────────────────────────────────────────
@app.route('/debug/env')
@admin_required
def debug_env():
    """Ruta temporal de diagnostico. Borrar despues de resolver el problema de login."""
    def mask(val, keep_start=12, keep_end=8):
        if not val:
            return 'NO CONFIGURADA (vacia o None)'
        val = str(val)
        if len(val) <= keep_start + keep_end:
            return f'(muy corta, {len(val)} caracteres) {val[:3]}...'
        return f'{val[:keep_start]}...{val[-keep_end:]}  (longitud total: {len(val)})'

    client_id = os.environ.get('GOOGLE_CLIENT_ID')
    client_secret = os.environ.get('GOOGLE_CLIENT_SECRET')
    secret_key = os.environ.get('FLASK_SECRET_KEY')
    admin_email = os.environ.get('ADMIN_EMAIL')
    db_url = os.environ.get('DATABASE_URL')

    db_status = 'NO CONFIGURADA — la app no puede guardar datos de forma persistente sin esto'
    db_conn_test = 'N/A'
    if db_url:
        db_status = mask(db_url, keep_start=15, keep_end=10)
        try:
            test_conn = get_db()
            test_conn.execute('SELECT 1')
            test_conn.close()
            db_conn_test = 'OK — se pudo conectar y ejecutar una consulta de prueba'
        except Exception as e:
            db_conn_test = f'ERROR al conectar: {e}'

    lines = [
        f"DATABASE_URL: {db_status}",
        f"  -> prueba de conexion: {db_conn_test}",
        "",
        f"GOOGLE_CLIENT_ID: {mask(client_id)}",
        f"  -> termina en .apps.googleusercontent.com: {str(client_id).strip().endswith('.apps.googleusercontent.com') if client_id else 'N/A'}",
        f"  -> tiene espacios al inicio/final sin recortar: {(client_id != client_id.strip()) if client_id else 'N/A'}",
        "",
        f"GOOGLE_CLIENT_SECRET: {'CONFIGURADA (longitud ' + str(len(client_secret)) + ')' if client_secret else 'NO CONFIGURADA'}",
        "",
        f"FLASK_SECRET_KEY: {'CONFIGURADA' if secret_key else 'NO CONFIGURADA (usando valor por defecto, inseguro)'}",
        "",
        f"ADMIN_EMAIL: {admin_email if admin_email else 'NO CONFIGURADA'}",
        "",
        f"URL de callback que la app va a pedirle a Google: {url_for('auth_callback', _external=True)}",
    ]
    return "<pre style='font-family:monospace;font-size:14px;padding:20px'>" + "\n".join(lines) + "</pre>"

@app.route('/login')
def login_page():
    if 'user_email' in session:
        return redirect('/')
    return render_template('login.html')

@app.route('/auth/google')
def auth_google():
    redirect_uri = url_for('auth_callback', _external=True)
    return google.authorize_redirect(redirect_uri)

@app.route('/auth/callback')
def auth_callback():
    token = google.authorize_access_token()
    user_info = token.get('userinfo')
    if not user_info or not user_info.get('email'):
        return render_template('login.html', error='No se pudo verificar tu cuenta de Google. Intentá de nuevo.')

    email = user_info['email'].strip().lower()
    conn = get_db()
    user = conn.execute('SELECT * FROM usuarios WHERE email=?', (email,)).fetchone()
    conn.close()

    if not user or not user['active']:
        return render_template('login.html',
            error=f'La cuenta {email} no tiene acceso autorizado. Pedile a un administrador que te agregue.')

    session['user_email'] = email
    session['user_name'] = user_info.get('name', email)
    session['user_picture'] = user_info.get('picture', '')
    session['user_role'] = user['role']
    return redirect('/')

@app.route('/logout')
def logout():
    session.clear()
    return redirect('/login')

# ── Routes ────────────────────────────────────────────────────────────────
@app.route('/')
@login_required
def index():
    return render_template('index.html',
        is_admin=(session.get('user_role')=='admin'),
        user_name=session.get('user_name',''),
        user_email=session.get('user_email',''),
        user_picture=session.get('user_picture',''))

# ── Usuarios (solo admin) ──
@app.route('/api/usuarios', methods=['GET'])
@admin_required
def get_usuarios():
    conn = get_db()
    rows = conn.execute('SELECT * FROM usuarios ORDER BY email').fetchall()
    locales_rows = conn.execute('SELECT usuario_id, local_id FROM usuario_locales').fetchall()
    conn.close()
    locales_por_usuario = {}
    for lr in locales_rows:
        locales_por_usuario.setdefault(lr['usuario_id'], []).append(lr['local_id'])
    return jsonify([{'id': r['id'], 'email': r['email'], 'nombre': r['nombre'],
                     'role': r['role'], 'active': bool(r['active']),
                     'locales': locales_por_usuario.get(r['id'], [])} for r in rows])

def _set_usuario_locales(conn, usuario_id, local_ids):
    conn.execute('DELETE FROM usuario_locales WHERE usuario_id=?', (usuario_id,))
    for lid in (local_ids or []):
        conn.execute('INSERT INTO usuario_locales (usuario_id, local_id) VALUES (?,?) ON CONFLICT DO NOTHING', (usuario_id, lid))

@app.route('/api/usuarios', methods=['POST'])
@admin_required
def create_usuario():
    data = request.json
    email = data['email'].strip().lower()
    conn = get_db()
    try:
        cur = conn.execute('INSERT INTO usuarios (email, nombre, role, active, created_at) VALUES (?,?,?,?,?) RETURNING id',
                     (email, data.get('nombre',''), data.get('role','user'),
                      int(data.get('active', True)), datetime.now().isoformat()))
        new_id = cur.fetchone()['id']
        _set_usuario_locales(conn, new_id, data.get('locales', []))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese email ya está registrado'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/usuarios/<int:usuario_id>', methods=['PUT'])
@admin_required
def update_usuario(usuario_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE usuarios SET email=?, nombre=?, role=?, active=? WHERE id=?',
                     (data['email'].strip().lower(), data.get('nombre',''), data.get('role','user'),
                      int(data.get('active', True)), usuario_id))
        _set_usuario_locales(conn, usuario_id, data.get('locales', []))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese email ya está registrado en otro usuario'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/usuarios/<int:usuario_id>', methods=['DELETE'])
@admin_required
def delete_usuario(usuario_id):
    conn = get_db()
    # Evitar que el admin se borre a si mismo y se quede afuera
    row = conn.execute('SELECT email FROM usuarios WHERE id=?', (usuario_id,)).fetchone()
    if row and row['email'] == session.get('user_email'):
        conn.close()
        return jsonify({'error': 'No podés eliminar tu propio usuario mientras estás conectado'}), 400
    conn.execute('DELETE FROM usuario_locales WHERE usuario_id=?', (usuario_id,))
    conn.execute('DELETE FROM usuarios WHERE id=?', (usuario_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Novedades (anuncios del admin a todos los usuarios, con lectura por usuario) ──
@app.route('/api/novedades', methods=['GET'])
@login_required
def get_novedades():
    conn = get_db()
    rows = conn.execute('SELECT * FROM novedades ORDER BY created_at DESC').fetchall()
    leidas = {r['novedad_id'] for r in conn.execute(
        'SELECT novedad_id FROM novedades_leidas WHERE user_email=?', (session['user_email'],)).fetchall()}
    conn.close()
    return jsonify([{'id': r['id'], 'titulo': r['titulo'], 'desarrollo': r['desarrollo'],
                     'creado_por_nombre': r['creado_por_nombre'],
                     'created_at': r['created_at'].isoformat() if r['created_at'] else None,
                     'leida': r['id'] in leidas} for r in rows])

@app.route('/api/novedades/pendientes-count', methods=['GET'])
@login_required
def get_novedades_pendientes_count():
    conn = get_db()
    count = conn.execute('''
        SELECT COUNT(*) AS cnt FROM novedades n
        WHERE NOT EXISTS (
            SELECT 1 FROM novedades_leidas nl WHERE nl.novedad_id = n.id AND nl.user_email = ?
        )
    ''', (session['user_email'],)).fetchone()['cnt']
    conn.close()
    return jsonify({'count': count})

@app.route('/api/novedades', methods=['POST'])
@admin_required
def create_novedad():
    data = request.json
    titulo = (data.get('titulo') or '').strip()
    desarrollo = (data.get('desarrollo') or '').strip()
    if not titulo or not desarrollo:
        return jsonify({'error': 'Completá el título y el desarrollo'}), 400
    conn = get_db()
    conn.execute('INSERT INTO novedades (titulo, desarrollo, creado_por, creado_por_nombre) VALUES (?,?,?,?)',
                 (titulo, desarrollo, session['user_email'], session.get('user_name')))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/novedades/<int:nov_id>', methods=['PUT'])
@admin_required
def update_novedad(nov_id):
    data = request.json
    titulo = (data.get('titulo') or '').strip()
    desarrollo = (data.get('desarrollo') or '').strip()
    if not titulo or not desarrollo:
        return jsonify({'error': 'Completá el título y el desarrollo'}), 400
    conn = get_db()
    conn.execute('UPDATE novedades SET titulo=?, desarrollo=? WHERE id=?', (titulo, desarrollo, nov_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/novedades/<int:nov_id>', methods=['DELETE'])
@admin_required
def delete_novedad(nov_id):
    conn = get_db()
    conn.execute('DELETE FROM novedades_leidas WHERE novedad_id=?', (nov_id,))
    conn.execute('DELETE FROM novedades WHERE id=?', (nov_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/novedades/<int:nov_id>/marcar-leida', methods=['POST'])
@login_required
def marcar_novedad_leida(nov_id):
    conn = get_db()
    conn.execute('INSERT INTO novedades_leidas (novedad_id, user_email) VALUES (?,?) ON CONFLICT (novedad_id, user_email) DO NOTHING',
                 (nov_id, session['user_email']))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Sugerencias (mejoras, correcciones de recetas, nuevas herramientas) ──
@app.route('/api/sugerencias', methods=['GET'])
@login_required
def get_sugerencias():
    conn = get_db()
    if session.get('user_role') == 'admin':
        rows = conn.execute('SELECT * FROM sugerencias ORDER BY created_at DESC').fetchall()
    else:
        rows = conn.execute('SELECT * FROM sugerencias WHERE creado_por=? ORDER BY created_at DESC',
                             (session['user_email'],)).fetchall()
        # Al entrar a ver su propia lista, el usuario ya "vio" las respuestas/cambios de estado
        conn.execute('UPDATE sugerencias SET visto=TRUE WHERE creado_por=? AND visto=FALSE',
                     (session['user_email'],))
        conn.commit()
    conn.close()
    return jsonify([{'id': r['id'], 'tipo': r['tipo'], 'titulo': r['titulo'], 'descripcion': r['descripcion'],
                     'estado': r['estado'], 'respuesta_admin': r['respuesta_admin'],
                     'creado_por': r['creado_por'], 'creado_por_nombre': r['creado_por_nombre'],
                     'referencia_tipo': r['referencia_tipo'], 'referencia_id': r['referencia_id'],
                     'referencia_nombre': r['referencia_nombre'],
                     'created_at': str(r['created_at']), 'updated_at': str(r['updated_at'])} for r in rows])

@app.route('/api/sugerencias/pendientes-count', methods=['GET'])
@login_required
def get_sugerencias_pendientes_count():
    conn = get_db()
    if session.get('user_role') == 'admin':
        count = conn.execute("SELECT COUNT(*) AS cnt FROM sugerencias WHERE estado='nueva'").fetchone()['cnt']
    else:
        count = conn.execute('SELECT COUNT(*) AS cnt FROM sugerencias WHERE creado_por=? AND visto=FALSE',
                             (session['user_email'],)).fetchone()['cnt']
    conn.close()
    return jsonify({'count': count})

@app.route('/api/sugerencias', methods=['POST'])
@login_required
def create_sugerencia():
    data = request.json
    titulo = (data.get('titulo') or '').strip()
    descripcion = (data.get('descripcion') or '').strip()
    tipo = data.get('tipo')
    if not titulo or not descripcion or tipo not in ('mejora', 'correccion_receta', 'nueva_herramienta'):
        return jsonify({'error': 'Completá el tipo, el título y la descripción'}), 400
    ref_tipo = data.get('referencia_tipo')
    if ref_tipo not in ('roll', 'semielaborado', 'combo', 'otro_producto'):
        ref_tipo = None
    conn = get_db()
    conn.execute('''INSERT INTO sugerencias (tipo, titulo, descripcion, creado_por, creado_por_nombre,
                    referencia_tipo, referencia_id, referencia_nombre)
                    VALUES (?,?,?,?,?,?,?,?)''',
                 (tipo, titulo, descripcion, session['user_email'], session.get('user_name', ''),
                  ref_tipo, data.get('referencia_id') if ref_tipo else None,
                  data.get('referencia_nombre') if ref_tipo else None))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/sugerencias/<int:sug_id>', methods=['PUT'])
@admin_required
def update_sugerencia(sug_id):
    data = request.json
    estado = data.get('estado')
    if estado not in ('nueva', 'en_revision', 'aceptada', 'rechazada', 'implementada'):
        return jsonify({'error': 'Estado inválido'}), 400
    conn = get_db()
    conn.execute('''UPDATE sugerencias SET estado=?, respuesta_admin=?, visto=FALSE, updated_at=CURRENT_TIMESTAMP WHERE id=?''',
                 (estado, data.get('respuesta_admin') or None, sug_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/sugerencias/<int:sug_id>', methods=['DELETE'])
@login_required
def delete_sugerencia(sug_id):
    conn = get_db()
    row = conn.execute('SELECT creado_por FROM sugerencias WHERE id=?', (sug_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'error': 'No encontrada'}), 404
    if session.get('user_role') != 'admin' and row['creado_por'] != session['user_email']:
        conn.close()
        return jsonify({'error': 'No podés borrar una sugerencia de otro usuario'}), 403
    conn.execute('DELETE FROM sugerencias WHERE id=?', (sug_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Locales ──
@app.route('/api/locales', methods=['GET'])
@login_required
def get_locales():
    allowed = get_user_local_ids()
    conn = get_db()
    if allowed is None:
        rows = conn.execute('SELECT * FROM locales ORDER BY name').fetchall()
    elif not allowed:
        rows = []
    else:
        placeholders = ','.join('?' * len(allowed))
        rows = conn.execute(f'SELECT * FROM locales WHERE id IN ({placeholders}) ORDER BY name', tuple(allowed)).fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'], 'active': bool(r['active']),
                     'marcas': json.loads(r['marcas'] or '[]')} for r in rows])

@app.route('/api/locales', methods=['POST'])
@admin_required
def create_local():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO locales (name, active, marcas) VALUES (?,?,?)',
                     (data['name'].strip(), int(data.get('active', True)), json.dumps(data.get('marcas', []))))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un local con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/locales/<int:local_id>', methods=['PUT'])
@admin_required
def update_local(local_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE locales SET name=?, active=?, marcas=? WHERE id=?',
                     (data['name'].strip(), int(data.get('active', True)), json.dumps(data.get('marcas', [])), local_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro local con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/locales/<int:local_id>', methods=['DELETE'])
@admin_required
def delete_local(local_id):
    conn = get_db()
    n_planillas = conn.execute('SELECT COUNT(*) AS cnt FROM planillas WHERE local_id=?', (local_id,)).fetchone()['cnt']
    if n_planillas > 0:
        conn.close()
        return jsonify({'error': f'Este local tiene {n_planillas} planilla(s) guardadas — no se puede eliminar. Podés desactivarlo en cambio.'}), 400
    conn.execute('DELETE FROM usuario_locales WHERE local_id=?', (local_id,))
    conn.execute('UPDATE sushimanes SET local_id=NULL WHERE local_id=?', (local_id,))
    conn.execute('DELETE FROM locales WHERE id=?', (local_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Respaldo completo (backup / restore) — solo admin ──────────────────────
@app.route('/api/backup', methods=['GET'])
@admin_required
def backup_data():
    conn = get_db()
    TABLES = ['combos', 'rolls', 'semielaborados', 'sushimanes', 'insumos', 'marcas', 'usuarios', 'familias', 'otros_productos', 'equivalencias', 'categorias_insumos', 'zonas_almacenamiento', 'proveedores', 'rollo_blanco_grupos', 'locales', 'usuario_locales', 'unidades', 'posiciones']
    data = {'version': 1, 'exported_at': datetime.now().isoformat()}
    for table in TABLES:
        rows = conn.execute(f'SELECT * FROM {table}').fetchall()
        data[table] = [dict(r) for r in rows]
    conn.close()

    resp = jsonify(data)
    filename = f'backup_sushi_{datetime.now().strftime("%Y%m%d_%H%M")}.json'
    resp.headers['Content-Disposition'] = f'attachment; filename={filename}'
    return resp

@app.route('/api/restore', methods=['POST'])
@admin_required
def restore_data():
    data = request.json
    if not data or 'version' not in data:
        return jsonify({'error': 'El archivo no parece ser un respaldo válido'}), 400

    conn = get_db()
    counts = {}

    def upsert(table, rows, unique_col, insert_cols):
        n = 0
        for row in rows:
            key_val = row.get(unique_col)
            if key_val is None:
                continue
            existing = conn.execute(f'SELECT id FROM {table} WHERE {unique_col}=?', (key_val,)).fetchone()
            values = [row.get(c) for c in insert_cols]
            if existing:
                set_clause = ', '.join(f'{c}=?' for c in insert_cols)
                conn.execute(f'UPDATE {table} SET {set_clause} WHERE id=?', values + [existing['id']])
            else:
                cols_clause = ', '.join(insert_cols)
                placeholders = ', '.join('?' for _ in insert_cols)
                conn.execute(f'INSERT INTO {table} ({cols_clause}) VALUES ({placeholders})', values)
            n += 1
        return n

    counts['combos'] = upsert('combos', data.get('combos', []), 'name',
                               ['name', 'family', 'rolls', 'marcas'])
    # Backups viejos (de antes de agregar "piezas por rollo") no tienen este campo;
    # les asignamos 14 por defecto para no romper la restauracion.
    for r in data.get('rolls', []):
        if r.get('piezas_por_rollo') is None:
            r['piezas_por_rollo'] = 14
        if r.get('cuenta_productividad') is None:
            r['cuenta_productividad'] = True
        if r.get('toppings') is None:
            r['toppings'] = '[]'
    counts['rolls'] = upsert('rolls', data.get('rolls', []), 'name',
                              ['name', 'insumos', 'marcas', 'piezas_por_rollo', 'rollo_blanco_grupo', 'cuenta_productividad', 'toppings'])
    counts['semielaborados'] = upsert('semielaborados', data.get('semielaborados', []), 'name',
                              ['name', 'insumo_key', 'unit', 'rolls', 'receta',
                               'rendimiento_cantidad', 'rendimiento_unidad', 'marcas',
                               'tiempo_elaboracion_min', 'vida_util_dias'])
    for sm in data.get('sushimanes', []):
        if sm.get('dias_franco') is None:
            sm['dias_franco'] = '[]'
        elif isinstance(sm['dias_franco'], list):
            sm['dias_franco'] = json.dumps(sm['dias_franco'])
    counts['sushimanes'] = upsert('sushimanes', data.get('sushimanes', []), 'name',
                              ['name', 'productivity', 'active', 'dias_franco', 'horario_ingreso', 'turno', 'posicion_id', 'local_id'])
    for p in data.get('posiciones', []):
        if p.get('orden') is None:
            p['orden'] = 0
    counts['posiciones'] = upsert('posiciones', data.get('posiciones', []), 'name', ['name', 'orden'])
    # Backups viejos no tienen los atributos nuevos de insumos; les damos defaults seguros
    for i in data.get('insumos', []):
        if i.get('es_80_20') is None:
            i['es_80_20'] = False
        if i.get('eficiencia') is None:
            i['eficiencia'] = 100
    counts['insumos'] = upsert('insumos', data.get('insumos', []), 'key',
                              ['key', 'label', 'unidad_receta', 'unidad_resumen',
                               'factor_conversion', 'precio_unidad', 'categoria', 'es_80_20',
                               'comentario', 'marca_producto', 'marca_tipo', 'zona_almacenamiento',
                               'proveedor_principal_id', 'proveedor_alt1_id', 'proveedor_alt2_id', 'eficiencia'])
    for m in data.get('marcas', []):
        if m.get('orden') is None:
            m['orden'] = 0
    counts['marcas'] = upsert('marcas', data.get('marcas', []), 'name', ['name', 'orden'])
    counts['usuarios'] = upsert('usuarios', data.get('usuarios', []), 'email',
                              ['email', 'nombre', 'role', 'active', 'created_at'])
    counts['familias'] = upsert('familias', data.get('familias', []), 'name', ['name'])
    for p in data.get('otros_productos', []):
        if p.get('rolls') is None:
            p['rolls'] = '{}'
        if p.get('procedimiento') is None:
            p['procedimiento'] = '[]'
    counts['otros_productos'] = upsert('otros_productos', data.get('otros_productos', []), 'name',
                              ['name', 'tipo', 'familia', 'insumos', 'marcas', 'rolls', 'procedimiento'])
    counts['categorias_insumos'] = upsert('categorias_insumos', data.get('categorias_insumos', []), 'name', ['name'])
    counts['zonas_almacenamiento'] = upsert('zonas_almacenamiento', data.get('zonas_almacenamiento', []), 'name', ['name'])
    for u in data.get('unidades', []):
        if u.get('orden') is None:
            u['orden'] = 0
    counts['unidades'] = upsert('unidades', data.get('unidades', []), 'simbolo', ['nombre', 'simbolo', 'orden'])
    counts['proveedores'] = upsert('proveedores', data.get('proveedores', []), 'name', ['name', 'contactos'])
    counts['rollo_blanco_grupos'] = upsert('rollo_blanco_grupos', data.get('rollo_blanco_grupos', []), 'name', ['name'])
    counts['locales'] = upsert('locales', data.get('locales', []), 'name', ['name', 'active', 'marcas'])

    # usuario_locales: clave compuesta, reemplazo completo simple (igual que equivalencias)
    ul_rows = data.get('usuario_locales', [])
    conn.execute('DELETE FROM usuario_locales')
    n_ul = 0
    for ul in ul_rows:
        if ul.get('usuario_id') is None or ul.get('local_id') is None:
            continue
        conn.execute('INSERT INTO usuario_locales (usuario_id, local_id) VALUES (?,?) ON CONFLICT DO NOTHING',
                     (ul['usuario_id'], ul['local_id']))
        n_ul += 1
    counts['usuario_locales'] = n_ul


    # Equivalencias: clave compuesta (tipo, nombre_alias), no calza con el helper 'upsert' generico.
    # Reemplazo completo simple: borramos todo y volvemos a insertar lo que venga en el backup.
    equiv_rows = data.get('equivalencias', [])
    conn.execute('DELETE FROM equivalencias')
    n_equiv = 0
    for e in equiv_rows:
        if not e.get('nombre_alias'):
            continue
        conn.execute('INSERT INTO equivalencias (tipo, nombre_canonico, nombre_alias, marca, factor) VALUES (?,?,?,?,?)',
                     (e['tipo'], e['nombre_canonico'], e['nombre_alias'], e.get('marca', ''), e.get('factor') or 1.0))
        n_equiv += 1
    counts['equivalencias'] = n_equiv

    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'counts': counts})

# ── Familias ──
@app.route('/api/familias', methods=['GET'])
@login_required
def get_familias():
    conn = get_db()
    rows = conn.execute('SELECT * FROM familias ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])

@app.route('/api/familias', methods=['POST'])
@admin_required
def create_familia():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO familias (name) VALUES (?)', (data['name'].strip(),))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Esa familia ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/familias/<int:familia_id>', methods=['PUT'])
@admin_required
def update_familia(familia_id):
    data = request.json
    conn = get_db()
    old = conn.execute('SELECT name FROM familias WHERE id=?', (familia_id,)).fetchone()
    new_name = data['name'].strip()
    try:
        conn.execute('UPDATE familias SET name=? WHERE id=?', (new_name, familia_id))
        # Si se renombra, actualizamos las referencias existentes en combos y otros_productos
        if old and old['name'] != new_name:
            conn.execute('UPDATE combos SET family=? WHERE family=?', (new_name, old['name']))
            conn.execute('UPDATE otros_productos SET familia=? WHERE familia=?', (new_name, old['name']))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otra familia con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/familias/<int:familia_id>', methods=['DELETE'])
@admin_required
def delete_familia(familia_id):
    conn = get_db()
    conn.execute('DELETE FROM familias WHERE id=?', (familia_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Otros productos (porciones, ensaladas, entradas, platos calientes) ──
@app.route('/api/otros-productos', methods=['GET'])
@login_required
def get_otros_productos():
    conn = get_db()
    rows = conn.execute('SELECT * FROM otros_productos ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'], 'tipo': r['tipo'],
                     'familia': r['familia'], 'insumos': json.loads(r['insumos']),
                     'marcas': json.loads(r['marcas'] or '[]'),
                     'rolls': json.loads(r['rolls'] or '{}'),
                     'procedimiento': json.loads(r['procedimiento'] or '[]')} for r in rows])

@app.route('/api/otros-productos', methods=['POST'])
@admin_required
def create_otro_producto():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO otros_productos (name, tipo, familia, insumos, marcas, rolls, procedimiento) VALUES (?,?,?,?,?,?,?)',
                     (data['name'], data.get('tipo','porcion'), data.get('familia',''),
                      json.dumps(data.get('insumos', {})), json.dumps(data.get('marcas', [])),
                      json.dumps(data.get('rolls', {})), json.dumps(data.get('procedimiento', []))))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un producto con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/otros-productos/<int:producto_id>', methods=['PUT'])
@admin_required
def update_otro_producto(producto_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE otros_productos SET name=?, tipo=?, familia=?, insumos=?, marcas=?, rolls=?, procedimiento=? WHERE id=?',
                     (data['name'], data.get('tipo','porcion'), data.get('familia',''),
                      json.dumps(data.get('insumos', {})), json.dumps(data.get('marcas', [])),
                      json.dumps(data.get('rolls', {})), json.dumps(data.get('procedimiento', [])), producto_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro producto con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/otros-productos/<int:producto_id>', methods=['DELETE'])
@admin_required
def delete_otro_producto(producto_id):
    conn = get_db()
    conn.execute('DELETE FROM otros_productos WHERE id=?', (producto_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Importar recetas desde PDF (fichas técnicas de otro sistema) ──
def _slugify_insumo(nombre):
    key = nombre.strip().lower()
    key = ''.join(c if c.isalnum() or c in ' -_' else ' ' for c in key)
    key = '_'.join(key.split())
    return key or 'insumo'

def _fmt_cant(n):
    """Redondea a 2 decimales pero sin mostrar decimales de más (0.75 se
    muestra como '0.75', 3.0 se muestra como '3') — antes se usaba round()
    sin decimales, que convertía 0.75 hojas en 1 hoja."""
    r = round(n, 2)
    if r == int(r):
        return str(int(r))
    return f"{r:g}"

@app.route('/api/importar-recetas-pdf', methods=['POST'])
@admin_required
def importar_recetas_pdf():
    """Parsea uno o más PDFs de fichas técnicas y devuelve una vista previa
    con los ingredientes ya emparejados (o no) contra el catálogo existente.
    NO guarda nada todavía — eso lo hace /api/confirmar-importacion-recetas."""
    if 'files' not in request.files:
        return jsonify({'error': 'No se recibió ningún archivo'}), 400
    files = request.files.getlist('files')
    if not files:
        return jsonify({'error': 'No se recibió ningún archivo'}), 400

    conn = get_db()
    insumos_rows = conn.execute('SELECT key, label FROM insumos').fetchall()
    semis_rows = conn.execute('SELECT insumo_key AS key, name AS label FROM semielaborados').fetchall()
    catalogo = [dict(r) for r in insumos_rows] + [dict(r) for r in semis_rows]

    resultados = []
    for f in files:
        tmp_path = os.path.join(tempfile.gettempdir(), f'import_receta_{os.urandom(6).hex()}.pdf')
        f.save(tmp_path)
        try:
            parsed = parse_receta_pdf(tmp_path)
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

        if 'error' in parsed:
            resultados.append({'archivo': f.filename, 'error': parsed['error']})
            continue

        ingredientes_anotados = []
        for ing in parsed['ingredientes']:
            match = match_insumo(ing['nombre'], catalogo)
            ingredientes_anotados.append({
                'nombre': ing['nombre'],
                'peso_neto': ing['peso_neto'],
                'formato': ing['formato'],
                'unidad_sugerida': ing['unidad'],
                'match_sugerido': match,  # {'key','label','score'} o None
                'es_semielaborado_sugerido': ing.get('es_semielaborado_sugerido', False),
            })

        # Verificar si ya existe un producto/semielaborado con este mismo nombre,
        # y si ese existente ya tiene una receta cargada (para no pisar datos buenos sin avisar)
        ya_existe = False
        ya_existe_tiene_receta = False
        if parsed['es_semielaborado']:
            existente = conn.execute('SELECT receta FROM semielaborados WHERE LOWER(name)=LOWER(?)', (parsed['titulo'],)).fetchone()
            if existente:
                ya_existe = True
                ya_existe_tiene_receta = bool(json.loads(existente['receta'] or '[]'))
        else:
            existente = conn.execute('SELECT insumos FROM otros_productos WHERE LOWER(name)=LOWER(?)', (parsed['titulo'],)).fetchone()
            if existente:
                ya_existe = True
                ya_existe_tiene_receta = bool(json.loads(existente['insumos'] or '{}'))

        resultados.append({
            'archivo': f.filename,
            'titulo': parsed['titulo'],
            'categoria': parsed['categoria'],
            'es_semielaborado': parsed['es_semielaborado'],
            'tipo_sugerido': parsed['tipo'],  # puede ser None si no se reconoció la categoria
            'rendimiento_cantidad': parsed.get('rendimiento_cantidad'),
            'rendimiento_unidad': parsed.get('rendimiento_unidad'),
            'vida_util_dias': parsed.get('vida_util_dias'),
            'ingredientes': ingredientes_anotados,
            'procedimiento_pasos': parsed['procedimiento_pasos'],
            'raciones': parsed['raciones'],
            'cantidad': parsed['cantidad'],
            'ya_existe': ya_existe,
            'ya_existe_tiene_receta': ya_existe_tiene_receta,
        })

    conn.close()
    return jsonify({'resultados': resultados})

@app.route('/api/confirmar-importacion-recetas', methods=['POST'])
@admin_required
def confirmar_importacion_recetas():
    """Recibe la lista de recetas ya revisadas/confirmadas por el usuario
    (con el tipo definitivo, familia, y cada ingrediente ya vinculado a un
    insumo o semielaborado existente, o marcado para crear uno nuevo) y las guarda.
    Si una receta ya existe (mismo nombre) y viene con actualizar_existente=true,
    actualiza ese registro en lugar de crear uno nuevo — así conserva su id y su
    insumo_key, y sigue bien vinculado desde cualquier otra receta que ya lo use."""
    data = request.json
    recetas = data.get('recetas', [])
    conn = get_db()
    creados_insumos = 0
    creados_semis_placeholder = 0
    creados_productos = 0
    actualizados = 0
    errores = []

    def resolver_ingredientes(ingredientes):
        """Devuelve una lista de (key, cantidad) resolviendo/creando insumos o
        semielaborados nuevos donde haga falta. Actualiza los contadores por closure."""
        nonlocal creados_insumos, creados_semis_placeholder
        resueltos = []
        for ing in ingredientes:
            if ing.get('excluir'):
                continue
            key = ing.get('insumo_key')
            if not key and ing.get('crear_nuevo'):
                nueva_label = (ing.get('nombre') or '').strip()
                key = _slugify_insumo(nueva_label)
                ya_existe_insumo = conn.execute('SELECT 1 FROM insumos WHERE key=?', (key,)).fetchone()
                ya_existe_semi = conn.execute('SELECT 1 FROM semielaborados WHERE insumo_key=?', (key,)).fetchone()
                if not ya_existe_insumo and not ya_existe_semi:
                    unidad = ing.get('unidad') or 'g'
                    unidad_resumen = 'kg' if unidad == 'g' else ('l' if unidad == 'ml' else unidad)
                    factor = 0.001 if unidad in ('g', 'ml') else 1
                    conn.execute('''INSERT INTO insumos (key,label,unidad_receta,unidad_resumen,factor_conversion,categoria)
                                    VALUES (?,?,?,?,?,?)''',
                                 (key, nueva_label, unidad, unidad_resumen, factor, 'Importado desde PDF'))
                    creados_insumos += 1
            elif not key and ing.get('crear_semi_nuevo'):
                nueva_label = (ing.get('nombre') or '').strip()
                key = _slugify_insumo(nueva_label)
                ya_existe_insumo = conn.execute('SELECT 1 FROM insumos WHERE key=?', (key,)).fetchone()
                ya_existe_semi = conn.execute('SELECT 1 FROM semielaborados WHERE insumo_key=?', (key,)).fetchone()
                if not ya_existe_insumo and not ya_existe_semi:
                    unidad = ing.get('unidad') or 'g'
                    unit = 'u' if unidad == 'u' else 'g'
                    # Semielaborado "placeholder": queda creado pero sin receta propia
                    # todavía — se puede completar más tarde a mano o re-importando su
                    # propia ficha (que lo va a actualizar en vez de duplicarlo).
                    conn.execute('''INSERT INTO semielaborados (name, insumo_key, unit, rolls, receta, rendimiento_cantidad, rendimiento_unidad, marcas)
                                    VALUES (?,?,?,?,?,?,?,?)''',
                                 (nueva_label, key, unit, json.dumps({}), json.dumps([]), 0, unidad, json.dumps([])))
                    creados_semis_placeholder += 1
            if key and ing.get('peso_neto') is not None:
                resueltos.append((key, float(ing['peso_neto'])))
        return resueltos

    for receta in recetas:
        nombre = (receta.get('nombre') or '').strip()
        if not nombre:
            errores.append('(sin nombre): se omitió')
            continue

        if receta.get('es_semielaborado'):
            existente = conn.execute('SELECT id, insumo_key FROM semielaborados WHERE LOWER(name)=LOWER(?)', (nombre,)).fetchone()
            if existente and not receta.get('actualizar_existente'):
                errores.append(f'"{nombre}": ya existe un semielaborado con ese nombre, se omitió (no se marcó para actualizar)')
                continue

            resueltos = resolver_ingredientes(receta.get('ingredientes', []))
            receta_lista = [{'key': k, 'cantidad': c} for k, c in resueltos]
            rend_unidad = receta.get('rendimiento_unidad') or 'g'
            unit = 'u' if rend_unidad == 'u' else 'g'

            if existente:
                # Actualiza el registro existente MANTENIENDO su id e insumo_key,
                # para que las recetas que ya lo usaban sigan bien vinculadas.
                conn.execute('''UPDATE semielaborados SET unit=?, receta=?, rendimiento_cantidad=?, rendimiento_unidad=?, vida_util_dias=? WHERE id=?''',
                             (unit, json.dumps(receta_lista), receta.get('rendimiento_cantidad') or 0,
                              rend_unidad, receta.get('vida_util_dias'), existente['id']))
                actualizados += 1
            else:
                insumo_key = _slugify_insumo(nombre)
                ya_existe_key = conn.execute('SELECT 1 FROM insumos WHERE key=?', (insumo_key,)).fetchone() or \
                                 conn.execute('SELECT 1 FROM semielaborados WHERE insumo_key=?', (insumo_key,)).fetchone()
                if ya_existe_key:
                    insumo_key = f'{insumo_key}_2'
                conn.execute('''INSERT INTO semielaborados
                                (name, insumo_key, unit, rolls, receta, rendimiento_cantidad, rendimiento_unidad, marcas, vida_util_dias)
                                VALUES (?,?,?,?,?,?,?,?,?)''',
                             (nombre, insumo_key, unit, json.dumps({}), json.dumps(receta_lista),
                              receta.get('rendimiento_cantidad') or 0, rend_unidad, json.dumps([]),
                              receta.get('vida_util_dias')))
                creados_productos += 1
            continue

        tipo = receta.get('tipo')
        if tipo not in ('porcion', 'ensalada', 'entrada', 'plato_caliente'):
            errores.append(f'"{nombre}": tipo inválido, se omitió')
            continue

        existente = conn.execute('SELECT id FROM otros_productos WHERE LOWER(name)=LOWER(?)', (nombre,)).fetchone()
        if existente and not receta.get('actualizar_existente'):
            errores.append(f'"{nombre}": ya existe un producto con ese nombre, se omitió (no se marcó para actualizar)')
            continue

        resueltos = resolver_ingredientes(receta.get('ingredientes', []))
        insumos_dict = {}
        for key, cantidad in resueltos:
            insumos_dict[key] = insumos_dict.get(key, 0) + cantidad

        if existente:
            conn.execute('''UPDATE otros_productos SET tipo=?, familia=?, insumos=?, procedimiento=? WHERE id=?''',
                         (tipo, receta.get('familia', ''), json.dumps(insumos_dict),
                          json.dumps(receta.get('procedimiento_pasos', [])), existente['id']))
            actualizados += 1
        else:
            conn.execute('''INSERT INTO otros_productos (name, tipo, familia, insumos, marcas, rolls, procedimiento)
                            VALUES (?,?,?,?,?,?,?)''',
                         (nombre, tipo, receta.get('familia', ''), json.dumps(insumos_dict),
                          json.dumps([]), json.dumps({}), json.dumps(receta.get('procedimiento_pasos', []))))
            creados_productos += 1

    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'creados_insumos': creados_insumos,
                     'creados_semis_placeholder': creados_semis_placeholder,
                     'creados_productos': creados_productos, 'actualizados': actualizados,
                     'errores': errores})

# ── Equivalencias (mismo producto, distinto nombre por marca) ──
@app.route('/api/equivalencias', methods=['GET'])
@login_required
def get_equivalencias():
    conn = get_db()
    rows = conn.execute('SELECT * FROM equivalencias ORDER BY tipo, nombre_canonico, nombre_alias').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'tipo': r['tipo'], 'nombre_canonico': r['nombre_canonico'],
                     'nombre_alias': r['nombre_alias'], 'marca': r['marca'],
                     'factor': r['factor'] if r['factor'] is not None else 1.0} for r in rows])

@app.route('/api/equivalencias', methods=['POST'])
@admin_required
def create_equivalencia():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO equivalencias (tipo, nombre_canonico, nombre_alias, marca, factor) VALUES (?,?,?,?,?)',
                     (data['tipo'], data['nombre_canonico'], data['nombre_alias'].strip(), data.get('marca', ''),
                      data.get('factor', 1.0) or 1.0))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese alias ya está usado para otro producto de ese tipo'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/equivalencias/<int:equiv_id>', methods=['PUT'])
@admin_required
def update_equivalencia(equiv_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE equivalencias SET tipo=?, nombre_canonico=?, nombre_alias=?, marca=?, factor=? WHERE id=?',
                     (data['tipo'], data['nombre_canonico'], data['nombre_alias'].strip(), data.get('marca', ''),
                      data.get('factor', 1.0) or 1.0, equiv_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese alias ya está usado para otro producto de ese tipo'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/equivalencias/<int:equiv_id>', methods=['DELETE'])
@admin_required
def delete_equivalencia(equiv_id):
    conn = get_db()
    conn.execute('DELETE FROM equivalencias WHERE id=?', (equiv_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/equivalencias/sync-producto', methods=['POST'])
@admin_required
def sync_equivalencias_producto():
    """Reemplaza los códigos (sin marca, con factor) de UN producto puntual,
    sin tocar los códigos de ningún otro producto ni los de la matriz por marca."""
    data = request.json
    tipo = data.get('tipo')
    nombre_canonico = data.get('nombre_canonico')
    codigos = data.get('codigos', [])  # [{codigo, factor}]
    if tipo not in ('combo', 'roll', 'otro_producto') or not nombre_canonico:
        return jsonify({'error': 'Datos inválidos'}), 400

    conn = get_db()
    conn.execute("""DELETE FROM equivalencias WHERE tipo=? AND nombre_canonico=?
                    AND (marca IS NULL OR marca='')""", (tipo, nombre_canonico))
    count = 0
    for c in codigos:
        codigo = (c.get('codigo') or '').strip()
        if not codigo:
            continue
        factor = c.get('factor')
        try:
            factor = float(factor) if factor not in (None, '') else 1.0
        except (TypeError, ValueError):
            factor = 1.0
        try:
            conn.execute('INSERT INTO equivalencias (tipo, nombre_canonico, nombre_alias, marca, factor) VALUES (?,?,?,?,?)',
                         (tipo, nombre_canonico, codigo, '', factor))
            count += 1
        except IntegrityError:
            conn.rollback()
            conn.close()
            return jsonify({'error': f'El código "{codigo}" ya está usado en otro producto de ese tipo'}), 400
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'count': count})

@app.route('/api/equivalencias/sync', methods=['POST'])
@admin_required
def sync_equivalencias():
    """Reemplaza TODAS las equivalencias de un tipo (combo/roll/otro_producto)
    con el estado actual de la matriz producto x marca que manda el frontend.
    Las celdas vacias simplemente no generan fila."""
    data = request.json
    tipo = data.get('tipo')
    entries = data.get('entries', [])  # [{nombre_canonico, marca, nombre_alias}]
    if tipo not in ('combo', 'roll', 'otro_producto'):
        return jsonify({'error': 'Tipo inválido'}), 400

    # Validamos ANTES de tocar la base: dos productos distintos no pueden usar el
    # mismo texto como nombre de venta (rompería el matcheo, no sabría a cuál va).
    # Si lo hiciéramos después de borrar lo viejo, un error acá perdería los datos
    # existentes sin guardar nada nuevo en su lugar.
    vistos = {}
    for e in entries:
        alias = (e.get('nombre_alias') or '').strip()
        if not alias:
            continue
        clave = alias.lower()
        if clave in vistos and vistos[clave] != e['nombre_canonico']:
            return jsonify({'error': f'"{alias}" está repetido en más de un producto (en "{vistos[clave]}" y en "{e["nombre_canonico"]}"). '
                                      f'Cada nombre de venta tiene que ser único — dejá la celda vacía si esa marca no lo vende, '
                                      f'en vez de escribir "no aplica" o similar.'}), 400
        vistos[clave] = e['nombre_canonico']

    conn = get_db()
    # Solo tocamos las filas CON marca (las de esta matriz) — las que no tienen marca
    # son códigos propios de cada producto (con o sin factor) y no se deben borrar acá.
    conn.execute("DELETE FROM equivalencias WHERE tipo=? AND marca IS NOT NULL AND marca != ''", (tipo,))
    count = 0
    for e in entries:
        alias = (e.get('nombre_alias') or '').strip()
        if not alias:
            continue
        conn.execute('INSERT INTO equivalencias (tipo, nombre_canonico, nombre_alias, marca, factor) VALUES (?,?,?,?,?)',
                     (tipo, e['nombre_canonico'], alias, e.get('marca', ''), 1.0))
        count += 1
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'count': count})

# ── Categorías de insumos ──
@app.route('/api/categorias-insumos', methods=['GET'])
@login_required
def get_categorias_insumos():
    conn = get_db()
    rows = conn.execute('SELECT * FROM categorias_insumos ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])

@app.route('/api/categorias-insumos', methods=['POST'])
@admin_required
def create_categoria_insumo():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO categorias_insumos (name) VALUES (?)', (data['name'].strip(),))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Esa categoría ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/categorias-insumos/<int:cat_id>', methods=['PUT'])
@admin_required
def update_categoria_insumo(cat_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE categorias_insumos SET name=? WHERE id=?', (data['name'].strip(), cat_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Esa categoría ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/categorias-insumos/<int:cat_id>', methods=['DELETE'])
@admin_required
def delete_categoria_insumo(cat_id):
    conn = get_db()
    conn.execute('DELETE FROM categorias_insumos WHERE id=?', (cat_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Zonas de almacenamiento ──
@app.route('/api/zonas-almacenamiento', methods=['GET'])
@login_required
def get_zonas_almacenamiento():
    conn = get_db()
    rows = conn.execute('SELECT * FROM zonas_almacenamiento ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])

@app.route('/api/zonas-almacenamiento', methods=['POST'])
@admin_required
def create_zona_almacenamiento():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO zonas_almacenamiento (name) VALUES (?)', (data['name'].strip(),))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Esa zona ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/zonas-almacenamiento/<int:zona_id>', methods=['PUT'])
@admin_required
def update_zona_almacenamiento(zona_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE zonas_almacenamiento SET name=? WHERE id=?', (data['name'].strip(), zona_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Esa zona ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/zonas-almacenamiento/<int:zona_id>', methods=['DELETE'])
@admin_required
def delete_zona_almacenamiento(zona_id):
    conn = get_db()
    conn.execute('DELETE FROM zonas_almacenamiento WHERE id=?', (zona_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Unidades (para los insumos) ──
@app.route('/api/unidades', methods=['GET'])
@login_required
def get_unidades():
    conn = get_db()
    rows = conn.execute('SELECT * FROM unidades ORDER BY orden, nombre').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'nombre': r['nombre'], 'simbolo': r['simbolo'], 'orden': r['orden']} for r in rows])

@app.route('/api/unidades', methods=['POST'])
@admin_required
def create_unidad():
    data = request.json
    conn = get_db()
    try:
        siguiente_orden = conn.execute('SELECT COALESCE(MAX(orden),0)+1 AS n FROM unidades').fetchone()['n']
        conn.execute('INSERT INTO unidades (nombre, simbolo, orden) VALUES (?,?,?)',
                     (data['nombre'].strip(), data['simbolo'].strip(), siguiente_orden))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe una unidad con ese símbolo'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/unidades/<int:unidad_id>', methods=['PUT'])
@admin_required
def update_unidad(unidad_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE unidades SET nombre=?, simbolo=? WHERE id=?',
                     (data['nombre'].strip(), data['simbolo'].strip(), unidad_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otra unidad con ese símbolo'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/unidades/<int:unidad_id>', methods=['DELETE'])
@admin_required
def delete_unidad(unidad_id):
    conn = get_db()
    conn.execute('DELETE FROM unidades WHERE id=?', (unidad_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/unidades/reorder', methods=['PUT'])
@admin_required
def reorder_unidades():
    data = request.json
    ids = data.get('ids', [])
    conn = get_db()
    for idx, unidad_id in enumerate(ids):
        conn.execute('UPDATE unidades SET orden=? WHERE id=?', (idx, unidad_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Proveedores (con contactos) ──
@app.route('/api/proveedores', methods=['GET'])
@login_required
def get_proveedores():
    conn = get_db()
    rows = conn.execute('SELECT * FROM proveedores ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'],
                     'contactos': json.loads(r['contactos'] or '[]')} for r in rows])

@app.route('/api/proveedores', methods=['POST'])
@admin_required
def create_proveedor():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO proveedores (name, contactos) VALUES (?,?)',
                     (data['name'].strip(), json.dumps(data.get('contactos', []))))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese proveedor ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/proveedores/<int:prov_id>', methods=['PUT'])
@admin_required
def update_proveedor(prov_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE proveedores SET name=?, contactos=? WHERE id=?',
                     (data['name'].strip(), json.dumps(data.get('contactos', [])), prov_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese proveedor ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/proveedores/<int:prov_id>', methods=['DELETE'])
@admin_required
def delete_proveedor(prov_id):
    conn = get_db()
    # Desvincular este proveedor de cualquier insumo que lo tenga asignado, para no dejar referencias rotas
    conn.execute('UPDATE insumos SET proveedor_principal_id=NULL WHERE proveedor_principal_id=?', (prov_id,))
    conn.execute('UPDATE insumos SET proveedor_alt1_id=NULL WHERE proveedor_alt1_id=?', (prov_id,))
    conn.execute('UPDATE insumos SET proveedor_alt2_id=NULL WHERE proveedor_alt2_id=?', (prov_id,))
    conn.execute('DELETE FROM proveedores WHERE id=?', (prov_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Grupos de rollo blanco (rolls que comparten el mismo relleno) ──
@app.route('/api/rollo-blanco-grupos', methods=['GET'])
@login_required
def get_rollo_blanco_grupos():
    conn = get_db()
    rows = conn.execute('SELECT * FROM rollo_blanco_grupos ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])

@app.route('/api/rollo-blanco-grupos', methods=['POST'])
@admin_required
def create_rollo_blanco_grupo():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO rollo_blanco_grupos (name) VALUES (?)', (data['name'].strip(),))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese grupo ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/rollo-blanco-grupos/<int:grupo_id>', methods=['PUT'])
@admin_required
def update_rollo_blanco_grupo(grupo_id):
    data = request.json
    conn = get_db()
    old = conn.execute('SELECT name FROM rollo_blanco_grupos WHERE id=?', (grupo_id,)).fetchone()
    new_name = data['name'].strip()
    try:
        conn.execute('UPDATE rollo_blanco_grupos SET name=? WHERE id=?', (new_name, grupo_id))
        if old and old['name'] != new_name:
            conn.execute('UPDATE rolls SET rollo_blanco_grupo=? WHERE rollo_blanco_grupo=?', (new_name, old['name']))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ese grupo ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/rollo-blanco-grupos/<int:grupo_id>', methods=['DELETE'])
@admin_required
def delete_rollo_blanco_grupo(grupo_id):
    conn = get_db()
    row = conn.execute('SELECT name FROM rollo_blanco_grupos WHERE id=?', (grupo_id,)).fetchone()
    if row:
        conn.execute('UPDATE rolls SET rollo_blanco_grupo=NULL WHERE rollo_blanco_grupo=?', (row['name'],))
    conn.execute('DELETE FROM rollo_blanco_grupos WHERE id=?', (grupo_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Rolls ──
@app.route('/api/rolls', methods=['GET'])
@login_required
def get_rolls():
    conn = get_db()
    rows = conn.execute('SELECT * FROM rolls ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'],
                     'insumos': json.loads(r['insumos']),
                     'marcas': json.loads(r['marcas'] or '[]'),
                     'piezas_por_rollo': r['piezas_por_rollo'],
                     'rollo_blanco_grupo': r['rollo_blanco_grupo'],
                     'cuenta_productividad': bool(r['cuenta_productividad']),
                     'toppings': json.loads(r['toppings'] or '[]')} for r in rows])

@app.route('/api/rolls', methods=['POST'])
@admin_required
def create_roll():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO rolls (name, insumos, marcas, piezas_por_rollo, rollo_blanco_grupo, cuenta_productividad, toppings) VALUES (?,?,?,?,?,?,?)',
                     (data['name'], json.dumps(data['insumos']), json.dumps(data.get('marcas', [])),
                      data.get('piezas_por_rollo', 14), data.get('rollo_blanco_grupo') or None,
                      bool(data.get('cuenta_productividad', True)), json.dumps(data.get('toppings', []))))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un roll con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/rolls/<int:roll_id>', methods=['PUT'])
@admin_required
def update_roll(roll_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE rolls SET name=?, insumos=?, marcas=?, piezas_por_rollo=?, rollo_blanco_grupo=?, cuenta_productividad=?, toppings=? WHERE id=?',
                     (data['name'], json.dumps(data['insumos']), json.dumps(data.get('marcas', [])),
                      data.get('piezas_por_rollo', 14), data.get('rollo_blanco_grupo') or None,
                      bool(data.get('cuenta_productividad', True)), json.dumps(data.get('toppings', [])), roll_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro roll con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/rolls/<int:roll_id>', methods=['DELETE'])
@admin_required
def delete_roll(roll_id):
    conn = get_db()
    conn.execute('DELETE FROM rolls WHERE id=?', (roll_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Combos ──
@app.route('/api/combos', methods=['GET'])
@login_required
def get_combos():
    conn = get_db()
    rows = conn.execute('SELECT * FROM combos ORDER BY family, name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'], 'family': r['family'],
                     'rolls': json.loads(r['rolls']),
                     'marcas': json.loads(r['marcas'] or '[]')} for r in rows])

@app.route('/api/combos', methods=['POST'])
@admin_required
def create_combo():
    data = request.json
    conn = get_db()
    try:
        conn.execute('INSERT INTO combos (name, family, rolls, marcas) VALUES (?,?,?,?)',
                     (data['name'], data['family'], json.dumps(data['rolls']), json.dumps(data.get('marcas', []))))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un combo con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/combos/<int:combo_id>', methods=['PUT'])
@admin_required
def update_combo(combo_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE combos SET name=?, family=?, rolls=?, marcas=? WHERE id=?',
                     (data['name'], data['family'], json.dumps(data['rolls']), json.dumps(data.get('marcas', [])), combo_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro combo con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/combos/<int:combo_id>', methods=['DELETE'])
@admin_required
def delete_combo(combo_id):
    conn = get_db()
    conn.execute('DELETE FROM combos WHERE id=?', (combo_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Semielaborados ──
@app.route('/api/semielaborados', methods=['GET'])
@login_required
def get_semielaborados():
    conn = get_db()
    rows = conn.execute('SELECT * FROM semielaborados ORDER BY name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'], 'insumo_key': r['insumo_key'],
                     'unit': r['unit'], 'rolls': json.loads(r['rolls']),
                     'receta': json.loads(r['receta'] or '[]'),
                     'rendimiento_cantidad': r['rendimiento_cantidad'],
                     'rendimiento_unidad': r['rendimiento_unidad'],
                     'marcas': json.loads(r['marcas'] or '[]'),
                     'tiempo_elaboracion_min': r['tiempo_elaboracion_min'],
                     'vida_util_dias': r['vida_util_dias']} for r in rows])

@app.route('/api/semielaborados', methods=['POST'])
@admin_required
def create_semi():
    data = request.json
    conn = get_db()
    try:
        conn.execute('''INSERT INTO semielaborados
                        (name, insumo_key, unit, rolls, receta, rendimiento_cantidad, rendimiento_unidad, marcas,
                         tiempo_elaboracion_min, vida_util_dias)
                        VALUES (?,?,?,?,?,?,?,?,?,?)''',
                     (data['name'], data['insumo_key'], data['unit'], json.dumps(data['rolls']),
                      json.dumps(data.get('receta', [])),
                      data.get('rendimiento_cantidad', 0),
                      data.get('rendimiento_unidad', 'g'),
                      json.dumps(data.get('marcas', [])),
                      data.get('tiempo_elaboracion_min') or None,
                      data.get('vida_util_dias') or None))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un semielaborado con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/semielaborados/<int:semi_id>', methods=['PUT'])
@admin_required
def update_semi(semi_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('''UPDATE semielaborados SET name=?, insumo_key=?, unit=?, rolls=?,
                        receta=?, rendimiento_cantidad=?, rendimiento_unidad=?, marcas=?,
                        tiempo_elaboracion_min=?, vida_util_dias=? WHERE id=?''',
                     (data['name'], data['insumo_key'], data['unit'], json.dumps(data['rolls']),
                      json.dumps(data.get('receta', [])),
                      data.get('rendimiento_cantidad', 0),
                      data.get('rendimiento_unidad', 'g'),
                      json.dumps(data.get('marcas', [])),
                      data.get('tiempo_elaboracion_min') or None,
                      data.get('vida_util_dias') or None, semi_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro semielaborado con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/semielaborados/<int:semi_id>', methods=['DELETE'])
@admin_required
def delete_semi(semi_id):
    conn = get_db()
    conn.execute('DELETE FROM semielaborados WHERE id=?', (semi_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Locales: helpers de acceso ──
def get_user_local_ids():
    """None = sin restriccion (admin, ve todos). Si no, lista de ids permitidos."""
    if session.get('user_role') == 'admin':
        return None
    conn = get_db()
    rows = conn.execute('''SELECT ul.local_id FROM usuario_locales ul
                            JOIN usuarios u ON u.id=ul.usuario_id
                            WHERE u.email=?''', (session['user_email'],)).fetchall()
    conn.close()
    return [r['local_id'] for r in rows]

def user_has_local_access(local_id):
    if session.get('user_role') == 'admin':
        return True
    if local_id is None:
        return False
    allowed = get_user_local_ids()
    return local_id in (allowed or [])

# ── Sushimanes ──
@app.route('/api/sushimanes', methods=['GET'])
@login_required
def get_sushimanes():
    allowed = get_user_local_ids()
    conn = get_db()
    if allowed is None:
        rows = conn.execute('SELECT * FROM sushimanes ORDER BY name').fetchall()
    elif not allowed:
        rows = []
    else:
        placeholders = ','.join('?' * len(allowed))
        rows = conn.execute(f'SELECT * FROM sushimanes WHERE local_id IN ({placeholders}) ORDER BY name', tuple(allowed)).fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'],
                     'productivity': r['productivity'],
                     'active': bool(r['active']),
                     'dias_franco': json.loads(r['dias_franco'] or '[]'),
                     'horario_ingreso': r['horario_ingreso'],
                     'turno': r['turno'],
                     'posicion_id': r['posicion_id'],
                     'local_id': r['local_id']} for r in rows])

@app.route('/api/sushimanes', methods=['POST'])
@login_required
def create_sushiman():
    data = request.json
    local_id = data.get('local_id')
    if not local_id or not user_has_local_access(local_id):
        return jsonify({'error': 'No tenés acceso a ese local'}), 403
    conn = get_db()
    try:
        conn.execute('INSERT INTO sushimanes (name, productivity, dias_franco, horario_ingreso, turno, posicion_id, local_id) VALUES (?,?,?,?,?,?,?)',
                     (data['name'], data.get('productivity', 10),
                      json.dumps(data.get('dias_franco', [])), data.get('horario_ingreso') or None,
                      data.get('turno') or None, data.get('posicion_id') or None, local_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un sushiman con ese nombre en ese local'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/sushimanes/<int:sm_id>', methods=['PUT'])
@login_required
def update_sushiman(sm_id):
    data = request.json
    conn = get_db()
    existing = conn.execute('SELECT local_id FROM sushimanes WHERE id=?', (sm_id,)).fetchone()
    if not existing or not user_has_local_access(existing['local_id']):
        conn.close()
        return jsonify({'error': 'No tenés acceso a ese local'}), 403
    try:
        conn.execute('UPDATE sushimanes SET name=?, productivity=?, active=?, dias_franco=?, horario_ingreso=?, turno=?, posicion_id=? WHERE id=?',
                     (data['name'], data['productivity'], int(data.get('active', True)),
                      json.dumps(data.get('dias_franco', [])), data.get('horario_ingreso') or None,
                      data.get('turno') or None, data.get('posicion_id') or None, sm_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro sushiman con ese nombre en ese local'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/sushimanes/<int:sm_id>', methods=['DELETE'])
@login_required
def delete_sushiman(sm_id):
    conn = get_db()
    existing = conn.execute('SELECT local_id FROM sushimanes WHERE id=?', (sm_id,)).fetchone()
    if not existing or not user_has_local_access(existing['local_id']):
        conn.close()
        return jsonify({'error': 'No tenés acceso a ese local'}), 403
    conn.execute('DELETE FROM sushimanes WHERE id=?', (sm_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Posiciones (para los sushimanes) ──
@app.route('/api/posiciones', methods=['GET'])
@login_required
def get_posiciones():
    conn = get_db()
    rows = conn.execute('SELECT * FROM posiciones ORDER BY orden, name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'], 'orden': r['orden']} for r in rows])

@app.route('/api/posiciones', methods=['POST'])
@admin_required
def create_posicion():
    data = request.json
    conn = get_db()
    try:
        siguiente_orden = conn.execute('SELECT COALESCE(MAX(orden),0)+1 AS n FROM posiciones').fetchone()['n']
        conn.execute('INSERT INTO posiciones (name, orden) VALUES (?,?)', (data['name'].strip(), siguiente_orden))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe una posición con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/posiciones/<int:pos_id>', methods=['PUT'])
@admin_required
def update_posicion(pos_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE posiciones SET name=? WHERE id=?', (data['name'].strip(), pos_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otra posición con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/posiciones/<int:pos_id>', methods=['DELETE'])
@admin_required
def delete_posicion(pos_id):
    conn = get_db()
    conn.execute('UPDATE sushimanes SET posicion_id=NULL WHERE posicion_id=?', (pos_id,))
    conn.execute('DELETE FROM posiciones WHERE id=?', (pos_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/posiciones/reorder', methods=['PUT'])
@admin_required
def reorder_posiciones():
    data = request.json
    ids = data.get('ids', [])
    conn = get_db()
    for idx, pos_id in enumerate(ids):
        conn.execute('UPDATE posiciones SET orden=? WHERE id=?', (idx, pos_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Marcas ──
@app.route('/api/marcas', methods=['GET'])
@login_required
def get_marcas():
    conn = get_db()
    rows = conn.execute('SELECT * FROM marcas ORDER BY orden, name').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name'], 'orden': r['orden']} for r in rows])

@app.route('/api/marcas', methods=['POST'])
@admin_required
def create_marca():
    data = request.json
    conn = get_db()
    try:
        siguiente_orden = conn.execute('SELECT COALESCE(MAX(orden),0)+1 AS n FROM marcas').fetchone()['n']
        conn.execute('INSERT INTO marcas (name, orden) VALUES (?,?)', (data['name'], siguiente_orden))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Esa marca ya existe'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/marcas/reorder', methods=['PUT'])
@admin_required
def reorder_marcas():
    """Recibe la lista de ids de marca en el orden final deseado y reasigna 'orden' 0,1,2..."""
    data = request.json
    ids = data.get('ids', [])
    conn = get_db()
    for idx, marca_id in enumerate(ids):
        conn.execute('UPDATE marcas SET orden=? WHERE id=?', (idx, marca_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/marcas/<int:marca_id>', methods=['PUT'])
@admin_required
def update_marca(marca_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('UPDATE marcas SET name=? WHERE id=?', (data['name'], marca_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otra marca con ese nombre'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/marcas/<int:marca_id>', methods=['DELETE'])
@admin_required
def delete_marca(marca_id):
    conn = get_db()
    conn.execute('DELETE FROM marcas WHERE id=?', (marca_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Insumos (tabla maestra de unidades) ──
@app.route('/api/insumos', methods=['GET'])
@login_required
def get_insumos():
    conn = get_db()
    rows = conn.execute('SELECT * FROM insumos ORDER BY label').fetchall()
    conn.close()
    return jsonify([{
        'id': r['id'], 'key': r['key'], 'label': r['label'],
        'unidad_receta': r['unidad_receta'], 'unidad_resumen': r['unidad_resumen'],
        'factor_conversion': r['factor_conversion'], 'precio_unidad': r['precio_unidad'],
        'categoria': r['categoria'], 'es_80_20': bool(r['es_80_20']), 'comentario': r['comentario'],
        'marca_producto': r['marca_producto'], 'marca_tipo': r['marca_tipo'],
        'zona_almacenamiento': r['zona_almacenamiento'],
        'proveedor_principal_id': r['proveedor_principal_id'],
        'proveedor_alt1_id': r['proveedor_alt1_id'],
        'proveedor_alt2_id': r['proveedor_alt2_id'],
        'eficiencia': r['eficiencia'],
    } for r in rows])

@app.route('/api/insumos', methods=['POST'])
@admin_required
def create_insumo():
    data = request.json
    conn = get_db()
    try:
        conn.execute('''INSERT INTO insumos (key,label,unidad_receta,unidad_resumen,factor_conversion,precio_unidad,
                        categoria,es_80_20,comentario,marca_producto,marca_tipo,zona_almacenamiento,
                        proveedor_principal_id,proveedor_alt1_id,proveedor_alt2_id,eficiencia)
                        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                     (data['key'], data['label'], data['unidad_receta'], data['unidad_resumen'],
                      data['factor_conversion'], data.get('precio_unidad'),
                      data.get('categoria'), int(bool(data.get('es_80_20'))), data.get('comentario'),
                      data.get('marca_producto'), data.get('marca_tipo'), data.get('zona_almacenamiento'),
                      data.get('proveedor_principal_id'), data.get('proveedor_alt1_id'), data.get('proveedor_alt2_id'),
                      data.get('eficiencia', 100) or 100))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe un insumo con un nombre muy similar. Probá con un nombre distinto (ej: agregando la marca o una aclaración).'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/insumos/<int:ins_id>', methods=['PUT'])
@admin_required
def update_insumo(ins_id):
    data = request.json
    conn = get_db()
    try:
        conn.execute('''UPDATE insumos SET key=?, label=?, unidad_receta=?, unidad_resumen=?,
                        factor_conversion=?, precio_unidad=?, categoria=?, es_80_20=?, comentario=?,
                        marca_producto=?, marca_tipo=?, zona_almacenamiento=?,
                        proveedor_principal_id=?, proveedor_alt1_id=?, proveedor_alt2_id=?, eficiencia=? WHERE id=?''',
                     (data['key'], data['label'], data['unidad_receta'], data['unidad_resumen'],
                      data['factor_conversion'], data.get('precio_unidad'),
                      data.get('categoria'), int(bool(data.get('es_80_20'))), data.get('comentario'),
                      data.get('marca_producto'), data.get('marca_tipo'), data.get('zona_almacenamiento'),
                      data.get('proveedor_principal_id'), data.get('proveedor_alt1_id'), data.get('proveedor_alt2_id'),
                      data.get('eficiencia', 100) or 100,
                      ins_id))
        conn.commit()
    except IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify({'error': 'Ya existe otro insumo con un nombre muy similar. Probá con un nombre distinto.'}), 400
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/insumos/<int:ins_id>', methods=['DELETE'])
@admin_required
def delete_insumo(ins_id):
    conn = get_db()
    conn.execute('DELETE FROM insumos WHERE id=?', (ins_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

# ── Calcular producción ──
@app.route('/api/calcular', methods=['POST'])
@login_required
def calcular():
    body       = request.json
    sales      = body.get('sales', [])
    adj        = body.get('adjustments', {})
    adj_mode   = body.get('adjMode', 'pct')
    global_pct = body.get('globalPct', 100)

    conn = get_db()
    combos_db = {r['name']: json.loads(r['rolls'])
                 for r in conn.execute('SELECT name, rolls FROM combos').fetchall()}
    rolls_db  = {r['name']: json.loads(r['insumos'])
                 for r in conn.execute('SELECT name, insumos FROM rolls').fetchall()}
    roll_piezas_db = {r['name']: r['piezas_por_rollo']
                 for r in conn.execute('SELECT name, piezas_por_rollo FROM rolls').fetchall()}
    roll_grupo_db = {r['name']: r['rollo_blanco_grupo']
                 for r in conn.execute('SELECT name, rollo_blanco_grupo FROM rolls').fetchall()}
    roll_cuenta_prod_db = {r['name']: bool(r['cuenta_productividad'])
                 for r in conn.execute('SELECT name, cuenta_productividad FROM rolls').fetchall()}
    otros_db  = {r['name']: {'tipo': r['tipo'], 'insumos': json.loads(r['insumos']), 'rolls': json.loads(r['rolls'] or '{}')}
                 for r in conn.execute('SELECT name, tipo, insumos, rolls FROM otros_productos').fetchall()}
    semis_db  = conn.execute('SELECT * FROM semielaborados').fetchall()
    insumos_master = {r['key']: dict(r) for r in conn.execute('SELECT * FROM insumos').fetchall()}
    proveedores_db = {r['id']: r['name'] for r in conn.execute('SELECT id, name FROM proveedores').fetchall()}

    # Equivalencias: mismo producto, distinto nombre segun la marca que lo vendio
    def norm(s):
        return (s or '').lower().replace(' ', '').replace('-', '').replace('_', '')

    equiv_rows = conn.execute('SELECT tipo, nombre_canonico, nombre_alias, factor FROM equivalencias').fetchall()
    equiv_combo = {norm(r['nombre_alias']): (r['nombre_canonico'], r['factor'] or 1.0) for r in equiv_rows if r['tipo'] == 'combo'}
    equiv_otro  = {norm(r['nombre_alias']): (r['nombre_canonico'], r['factor'] or 1.0) for r in equiv_rows if r['tipo'] == 'otro_producto'}
    equiv_roll  = {norm(r['nombre_alias']): r['nombre_canonico'] for r in equiv_rows if r['tipo'] == 'roll'}
    conn.close()

    roll_totals = {}
    otros_totals = {}  # nombre -> {tipo, qty}
    direct_insumo_totals = {}  # insumos que vienen DIRECTO de "otros productos" (no de rolls)

    for s in sales:
        factor = 1.0
        combo = combos_db.get(s['name']) or combos_db.get(s['code'])
        if not combo:
            name_lower = s['name'].lower().replace(' ', '')
            for cname, crolls in combos_db.items():
                if cname.lower().replace(' ', '') == name_lower:
                    combo = crolls
                    break
        if not combo:
            # Ver si el nombre/codigo es un codigo/alias conocido para un combo existente
            equiv_match = equiv_combo.get(norm(s['name'])) or equiv_combo.get(norm(s['code']))
            if equiv_match:
                canonico, factor = equiv_match
                combo = combos_db.get(canonico)

        otro = None
        if not combo:
            # No es combo: ver si matchea con "otros productos" (porciones, ensaladas, entradas, platos calientes)
            otro = otros_db.get(s['name']) or otros_db.get(s['code'])
            if not otro:
                name_lower = s['name'].lower().replace(' ', '')
                for pname, pdata in otros_db.items():
                    if pname.lower().replace(' ', '') == name_lower:
                        otro = pdata
                        break
            if not otro:
                # Ver si el nombre/codigo es un codigo/alias conocido para un producto existente
                equiv_match = equiv_otro.get(norm(s['name'])) or equiv_otro.get(norm(s['code']))
                if equiv_match:
                    canonico, factor = equiv_match
                    otro = otros_db.get(canonico)

        if not combo and not otro:
            # No matchea ningun combo ni otro producto por ningun medio: se ignora por completo
            continue

        # El factor (si vino de un codigo con equivalencia, ej "UPS"=0.5) se aplica sobre
        # la cantidad base, ANTES de los ajustes por porcentaje/valor fijo y el % global.
        effective_qty = s['qty'] * factor
        a = adj.get(s['code'], 0)
        adj_qty = max(0, round(effective_qty + (effective_qty * a / 100 if adj_mode == 'pct' else a)))
        final_qty = round(adj_qty * global_pct / 100)
        if final_qty <= 0:
            continue

        if combo:
            for roll_name, piezas in combo.items():
                if not piezas:
                    continue
                roll_name = equiv_roll.get(norm(roll_name), roll_name)
                piezas_por_rollo = roll_piezas_db.get(roll_name, 14) or 14
                rolls_needed = -(-piezas * final_qty // piezas_por_rollo)
                roll_totals[roll_name] = roll_totals.get(roll_name, 0) + rolls_needed
            continue

        if otro:
            key = None
            for pname, pdata in otros_db.items():
                if pdata is otro:
                    key = pname
                    break
            if key:
                if key not in otros_totals:
                    otros_totals[key] = {'tipo': otro['tipo'], 'qty': 0}
                otros_totals[key]['qty'] += final_qty
            for k, v in otro['insumos'].items():
                if v:
                    direct_insumo_totals[k] = direct_insumo_totals.get(k, 0) + v * final_qty
            for roll_name, piezas in otro.get('rolls', {}).items():
                if not piezas:
                    continue
                roll_name = equiv_roll.get(norm(roll_name), roll_name)
                piezas_por_rollo = roll_piezas_db.get(roll_name, 14) or 14
                rolls_needed = -(-piezas * final_qty // piezas_por_rollo)
                roll_totals[roll_name] = roll_totals.get(roll_name, 0) + rolls_needed

    production = sorted(
        [{'name': k, 'qty': v, 'piezasPorRollo': roll_piezas_db.get(k, 14) or 14,
          'piezas': v * (roll_piezas_db.get(k, 14) or 14),
          'cuentaProductividad': roll_cuenta_prod_db.get(k, True)} for k, v in roll_totals.items() if v > 0],
        key=lambda x: -x['qty']
    )

    # Rollos blancos: agrupa los rolls que comparten el mismo relleno (misma base,
    # distinta cobertura) para poder armarlos todos juntos y despues diferenciar.
    grupos_blancos = {}
    for p in production:
        grupo = roll_grupo_db.get(p['name'])
        if grupo:
            grupos_blancos.setdefault(grupo, []).append({'name': p['name'], 'qty': p['qty']})
    rollos_blancos_out = sorted(
        [{'grupo': g, 'rolls': rolls, 'totalQty': sum(r['qty'] for r in rolls)}
         for g, rolls in grupos_blancos.items()],
        key=lambda x: -x['totalQty']
    )

    otros_productos_out = sorted(
        [{'name': k, 'tipo': v['tipo'], 'qty': v['qty']} for k, v in otros_totals.items() if v['qty'] > 0],
        key=lambda x: -x['qty']
    )

    # Insumos — usar tabla maestra para labels/unidades si existe
    FALLBACK_LABELS = {
        'salmon':'Salmón','queso':'Queso crema','palta':'Palta',
        'langos':'Langostinos rebozados','algas':'Algas','arroz':'Arroz',
        'grill':'Grill','tartar':'Tartar de salmón','kanikama':'Kanikama',
        'batata':'Hilos de batata','guac':'Guacamole','spicy':'Salsa spicy',
        'okinawa':'Manga okinawa','salmonCrispy':'Salmón crispy'
    }
    insumo_totals = {}
    semi_keys_set = {s['insumo_key'] for s in semis_db}  # se muestran aparte, en "Semielaborados"
    for r in production:
        recipe = rolls_db.get(r['name'], {})
        for k, v in recipe.items():
            if v and k not in semi_keys_set:
                insumo_totals[k] = insumo_totals.get(k, 0) + v * r['qty']
    # Sumar tambien lo que aportan directo los "otros productos"
    for k, v in direct_insumo_totals.items():
        if k not in semi_keys_set:
            insumo_totals[k] = insumo_totals.get(k, 0) + v

    insumos_out = {}
    for k, total_receta in sorted(insumo_totals.items(), key=lambda x: -x[1]):
        master = insumos_master.get(k)
        eficiencia = (master['eficiencia'] if master and master.get('eficiencia') else 100) or 100
        # Lo que pide la receta no es lo que hay que comprar: si hay merma (fruta/verdura
        # que se pudre, producto que queda pegado al envase, etc.), hay que comprar de más.
        total = total_receta / (eficiencia / 100) if eficiencia > 0 else total_receta
        nota_eficiencia = f" — incluye {round(100-eficiencia)}% de merma estimada" if eficiencia < 100 else ""
        if master:
            label = master['label']
            factor = master['factor_conversion']
            unidad_resumen = master['unidad_resumen']
            converted = total * factor
            display = f"{_fmt_cant(total)} {master['unidad_receta']} / {converted:.2f} {unidad_resumen}{nota_eficiencia}"
        else:
            label = FALLBACK_LABELS.get(k, k)
            unit = 'hojas' if k == 'algas' else ('u' if k == 'langos' else 'g')
            if unit == 'g':
                display = f"{_fmt_cant(total)} g / {total/1000:.2f} kg{nota_eficiencia}"
            elif unit == 'hojas':
                display = f"{total:.1f} hojas{nota_eficiencia}"
            else:
                display = f"{round(total)} u{nota_eficiencia}"
        insumos_out[k] = {
            'label': label,
            'total': round(total, 1) if total < 100 else round(total),
            'total_receta': round(total_receta, 1) if total_receta < 100 else round(total_receta),
            'eficiencia': eficiencia,
            'display': display
        }

    # Semielaborados — se detectan automáticamente según qué rolls producidos
    # (o "otros productos": porciones, ensaladas, entradas, platos calientes)
    # incluyen su insumo_key en su receta (no depende de una lista manual)
    semis_by_key = {s['insumo_key']: s['name'] for s in semis_db}
    semis_by_key_full = {s['insumo_key']: s for s in semis_db}

    # Para el reporte de "insumos totales" (neto de recetas + lo que se gasta
    # PREPARANDO cada semielaborado) hace falta bajar hasta los insumos crudos,
    # aunque un semielaborado use OTRO semielaborado en su propia receta.
    insumo_totals_expandido = dict(insumo_totals)
    insumos_sin_vincular = {}  # ingredientes de receta sin 'key' -> no se pueden ligar a un insumo real

    def expandir_insumos_de_semi(semi_key, cantidad_necesaria, visitados=frozenset()):
        if semi_key in visitados:
            return  # corta ante una dependencia circular entre semielaborados
        semi_row = semis_by_key_full.get(semi_key)
        if not semi_row:
            return
        receta = json.loads(semi_row['receta'] or '[]')
        rend_cant = semi_row['rendimiento_cantidad'] or 0
        if not receta or rend_cant <= 0:
            return
        scale = cantidad_necesaria / rend_cant
        for ing in receta:
            cant = ing['cantidad'] * scale
            k = ing.get('key')
            if k and insumos_master.get(k):
                insumo_totals_expandido[k] = insumo_totals_expandido.get(k, 0) + cant
            elif k and semis_by_key_full.get(k):
                expandir_insumos_de_semi(k, cant, visitados | {semi_key})
            else:
                nombre = ing.get('nombre') or k or '?'
                unidad = ing.get('unidad', 'g')
                clave = f"{nombre}|{unidad}"
                insumos_sin_vincular[clave] = insumos_sin_vincular.get(clave, 0) + cant

    def resolve_ing(ing):
        """Resuelve una fila de receta a {nombre, unidad, cantidad} sea cual sea
        el formato: nuevo (con 'key' hacia insumos/semielaborados) o viejo (texto libre)."""
        if ing.get('key'):
            k = ing['key']
            master = insumos_master.get(k)
            if master:
                return {'nombre': master['label'], 'unidad': master['unidad_receta'], 'cantidad': ing['cantidad']}
            semi_name = semis_by_key.get(k)
            if semi_name:
                return {'nombre': semi_name + ' (elaborado)', 'unidad': 'g', 'cantidad': ing['cantidad']}
            return {'nombre': k, 'unidad': ing.get('unidad', 'g'), 'cantidad': ing['cantidad']}
        return {'nombre': ing.get('nombre', '?'), 'unidad': ing.get('unidad', 'g'), 'cantidad': ing['cantidad']}

    semis_out = []
    for semi in semis_db:
        total = 0
        used_in = []
        for prod in production:
            recipe = rolls_db.get(prod['name'], {})
            amt = recipe.get(semi['insumo_key'], 0)
            if amt:
                total += amt * prod['qty']
                used_in.append(prod['name'])
        for otro in otros_productos_out:
            recipe = otros_db.get(otro['name'], {}).get('insumos', {})
            amt = recipe.get(semi['insumo_key'], 0)
            if amt:
                total += amt * otro['qty']
                used_in.append(otro['name'])
        if total > 0:
            expandir_insumos_de_semi(semi['insumo_key'], total)
            unit = semi['unit']
            display = (f"{round(total)} g / {total/1000:.2f} kg" if unit == 'g'
                       else f"{round(total)} u")

            semi_out = {
                'name': semi['name'],
                'usedIn': ', '.join(used_in),
                'display': display,
                'total': round(total)
            }

            # Si tiene receta cargada, calcular lotes e ingredientes crudos
            receta = json.loads(semi['receta'] or '[]')
            rend_cant = semi['rendimiento_cantidad'] or 0
            rend_unid = semi['rendimiento_unidad'] or unit
            if receta and rend_cant > 0:
                scale = total / rend_cant
                batches = -(-scale // 1)  # ceil
                semi_out['receta'] = {
                    'rendimiento_cantidad': rend_cant,
                    'rendimiento_unidad': rend_unid,
                    'lotes_necesarios': int(batches),
                    'escala_exacta': round(scale, 2),
                    'ingredientes': [
                        {
                            'nombre': resolve_ing(ing)['nombre'],
                            'cantidad_receta': ing['cantidad'],
                            'unidad': resolve_ing(ing)['unidad'],
                            'cantidad_total': round(ing['cantidad'] * scale, 1)
                        } for ing in receta
                    ]
                }
            semis_out.append(semi_out)

    # Reporte de "insumos totales": todo lo que hace falta comprar, sumando lo que
    # se usa directo mas lo que se gasta PREPARANDO cada semielaborado (ya bajado
    # a insumos crudos), con la categoria/proveedor de cada uno para poder pedirlo.
    insumos_totales_out = []
    for k, total_receta in sorted(insumo_totals_expandido.items(), key=lambda x: -x[1]):
        master = insumos_master.get(k)
        eficiencia = (master['eficiencia'] if master and master.get('eficiencia') else 100) or 100
        total_bruto = total_receta / (eficiencia / 100) if eficiencia > 0 else total_receta
        insumos_totales_out.append({
            'insumo': master['label'] if master else FALLBACK_LABELS.get(k, k),
            'cantidad_neta': round(total_receta, 2),
            'unidad': master['unidad_receta'] if master else 'g',
            'eficiencia': eficiencia,
            'cantidad_bruta': round(total_bruto, 2),
            'categoria': (master.get('categoria') if master else None) or '',
            'proveedor': (proveedores_db.get(master.get('proveedor_principal_id')) if master else None) or '',
        })
    for clave, total_receta in sorted(insumos_sin_vincular.items(), key=lambda x: -x[1]):
        nombre, unidad = clave.rsplit('|', 1)
        insumos_totales_out.append({
            'insumo': nombre + ' (sin vincular a un insumo)',
            'cantidad_neta': round(total_receta, 2),
            'unidad': unidad,
            'eficiencia': 100,
            'cantidad_bruta': round(total_receta, 2),
            'categoria': '',
            'proveedor': '',
        })

    return jsonify({
        'production': production,
        'insumos': insumos_out,
        'semis': semis_out,
        'otrosProductos': otros_productos_out,
        'rollosBlancos': rollos_blancos_out,
        'insumosTotales': insumos_totales_out
    })

# ── Generar PDF ──
# ── Planillas (historial de producción por local) ──
@app.route('/api/planillas', methods=['GET'])
@login_required
def get_planillas():
    allowed = get_user_local_ids()
    local_filter = request.args.get('local_id')
    conn = get_db()
    where = []
    params = []
    if allowed is not None:
        if not allowed:
            conn.close()
            return jsonify([])
        placeholders = ','.join('?' * len(allowed))
        where.append(f'local_id IN ({placeholders})')
        params.extend(allowed)
    if local_filter:
        where.append('local_id=?')
        params.append(int(local_filter))
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''
    rows = conn.execute(f'''SELECT id, local_id, fecha, estado, created_by, created_at, updated_at
                            FROM planillas {where_sql} ORDER BY updated_at DESC LIMIT 200''', tuple(params)).fetchall()
    locales = {l['id']: l['name'] for l in conn.execute('SELECT id, name FROM locales').fetchall()}
    conn.close()
    return jsonify([{'id': r['id'], 'local_id': r['local_id'], 'local_name': locales.get(r['local_id'], '?'),
                     'fecha': r['fecha'], 'estado': r['estado'], 'created_by': r['created_by'],
                     'created_at': str(r['created_at']), 'updated_at': str(r['updated_at'])} for r in rows])

@app.route('/api/planillas/<int:planilla_id>', methods=['GET'])
@login_required
def get_planilla(planilla_id):
    conn = get_db()
    row = conn.execute('SELECT * FROM planillas WHERE id=?', (planilla_id,)).fetchone()
    conn.close()
    if not row:
        return jsonify({'error': 'No encontrada'}), 404
    if not user_has_local_access(row['local_id']):
        return jsonify({'error': 'No tenés acceso a esa planilla'}), 403
    return jsonify({'id': row['id'], 'local_id': row['local_id'], 'fecha': row['fecha'],
                     'estado': row['estado'], 'data': json.loads(row['data'] or '{}'),
                     'created_by': row['created_by'], 'created_at': str(row['created_at']),
                     'updated_at': str(row['updated_at'])})

@app.route('/api/planillas', methods=['POST'])
@login_required
def create_planilla():
    data = request.json
    local_id = data.get('local_id')
    if not local_id or not user_has_local_access(local_id):
        return jsonify({'error': 'No tenés acceso a ese local'}), 403
    conn = get_db()
    cur = conn.execute('''INSERT INTO planillas (local_id, fecha, estado, data, created_by, updated_at)
                          VALUES (?,?,?,?,?,CURRENT_TIMESTAMP) RETURNING id''',
                 (local_id, data.get('fecha', datetime.now().strftime('%Y-%m-%d')),
                  data.get('estado', 'borrador'), json.dumps(data.get('data', {})),
                  session.get('user_email')))
    new_id = cur.fetchone()['id']
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'id': new_id})

@app.route('/api/planillas/<int:planilla_id>', methods=['PUT'])
@login_required
def update_planilla(planilla_id):
    data = request.json
    conn = get_db()
    existing = conn.execute('SELECT local_id FROM planillas WHERE id=?', (planilla_id,)).fetchone()
    if not existing:
        conn.close()
        return jsonify({'error': 'No encontrada'}), 404
    if not user_has_local_access(existing['local_id']):
        conn.close()
        return jsonify({'error': 'No tenés acceso a esa planilla'}), 403
    conn.execute('''UPDATE planillas SET fecha=?, estado=?, data=?, updated_at=CURRENT_TIMESTAMP WHERE id=?''',
                 (data.get('fecha'), data.get('estado', 'borrador'), json.dumps(data.get('data', {})), planilla_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/planillas/<int:planilla_id>', methods=['DELETE'])
@login_required
def delete_planilla(planilla_id):
    conn = get_db()
    existing = conn.execute('SELECT local_id FROM planillas WHERE id=?', (planilla_id,)).fetchone()
    if not existing:
        conn.close()
        return jsonify({'error': 'No encontrada'}), 404
    if not user_has_local_access(existing['local_id']):
        conn.close()
        return jsonify({'error': 'No tenés acceso a esa planilla'}), 403
    conn.execute('DELETE FROM planillas WHERE id=?', (planilla_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})

@app.route('/api/pdf', methods=['POST'])
@login_required
def generar_pdf():
    body = request.json
    tmp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    tmp.close()
    build_pdf(body, tmp.name)
    date_str = body.get('date', datetime.now().strftime('%d-%m-%Y'))
    return send_file(tmp.name, as_attachment=True,
                     download_name=f'produccion_{date_str}.pdf',
                     mimetype='application/pdf')

@app.route('/api/insumos-totales-xlsx', methods=['POST'])
@login_required
def generar_insumos_totales_xlsx():
    body = request.json
    rows = body.get('insumosTotales', [])
    date_str = body.get('date', datetime.now().strftime('%d-%m-%Y'))
    global_pct = body.get('globalPct', 100)
    tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    tmp.close()
    build_insumos_xlsx(rows, date_str, global_pct, tmp.name)
    return send_file(tmp.name, as_attachment=True,
                     download_name=f'insumos_totales_{date_str}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ── Importar recetas desde Excel ──────────────────────────────────────────
@app.route('/api/importar/rolls', methods=['POST'])
@admin_required
def importar_rolls():
    if 'file' not in request.files:
        return jsonify({'error': 'No se recibió archivo'}), 400
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'El archivo debe ser .xlsx o .xls'}), 400

    try:
        from openpyxl import load_workbook
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        file.save(tmp.name)
        tmp.close()

        wb = load_workbook(tmp.name, data_only=True)
        os.unlink(tmp.name)

        ws = None
        for name in wb.sheetnames:
            if 'ROLL' in name.upper():
                ws = wb[name]
                break
        if not ws:
            return jsonify({'error': 'No se encontró la solapa ROLLS en el archivo'}), 400

        headers = []
        for cell in ws[3]:
            v = cell.value
            headers.append(str(v).strip() if v else '')

        HEADER_MAP = {
            'salmón (g)':'salmon','salmon (g)':'salmon',
            'queso crema (g)':'queso','queso (g)':'queso',
            'palta (g)':'palta',
            'langostinos (u)':'langos','langostinos rebozados (u)':'langos',
            'algas (hojas)':'algas',
            'arroz (g)':'arroz',
            'grill (g)':'grill',
            'tartar salmón (g)':'tartar','tartar salmon (g)':'tartar','tartar (g)':'tartar',
            'kanikama (g)':'kanikama',
            'batata (g)':'batata','hilos de batata (g)':'batata',
            'guacamole (g)':'guac','guac (g)':'guac',
            'salsa spicy (g)':'spicy','spicy (g)':'spicy',
            'manga okinawa (g)':'okinawa','okinawa (g)':'okinawa',
            'atún (g)':'atun','atun (g)':'atun',
            'camarón (g)':'camaron','camaron (g)':'camaron',
            'mayonesa (g)':'mayonesa','ciboulette (g)':'ciboulette',
            'pepino (g)':'pepino','zanahoria (g)':'zanahoria',
            'pimiento (g)':'pimiento','cream cheese (g)':'cream_cheese',
            'nori extra (hojas)':'nori','sésamo (g)':'sesamo','sesamo (g)':'sesamo',
        }

        conn = get_db()
        creados, actualizados = [], []

        for row in ws.iter_rows(min_row=4, values_only=True):
            nombre = row[0] if row[0] else None
            if not nombre or str(nombre).strip() == '' or 'EJEMPLO' in str(nombre).upper():
                continue
            nombre = str(nombre).strip()

            piezas_por_rollo = 14
            if len(row) > 1 and row[1] is not None:
                try:
                    piezas_por_rollo = int(float(row[1]))
                except (ValueError, TypeError):
                    pass

            insumos = {}
            for ci, header in enumerate(headers[2:], 2):
                if ci >= len(row): break
                val = row[ci]
                if val is None: continue
                try: val = float(val)
                except: continue
                if val <= 0: continue
                h_lower = header.lower().strip()
                key = HEADER_MAP.get(h_lower, h_lower.replace(' ','_').replace('(','').replace(')','').strip('_'))
                if key: insumos[key] = val

            existing = conn.execute('SELECT id FROM rolls WHERE LOWER(name)=LOWER(?)', (nombre,)).fetchone()
            if existing:
                conn.execute('UPDATE rolls SET insumos=?, piezas_por_rollo=? WHERE id=?',
                             (json.dumps(insumos), piezas_por_rollo, existing['id']))
                actualizados.append(nombre)
            else:
                conn.execute('INSERT INTO rolls (name, insumos, piezas_por_rollo) VALUES (?,?,?)',
                             (nombre, json.dumps(insumos), piezas_por_rollo))
                creados.append(nombre)

        conn.commit()
        conn.close()
        return jsonify({'ok': True, 'creados': creados, 'actualizados': actualizados,
                        'total': len(creados) + len(actualizados)})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/importar/semielaborados', methods=['POST'])
@admin_required
def importar_semielaborados():
    if 'file' not in request.files:
        return jsonify({'error': 'No se recibió archivo'}), 400
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'El archivo debe ser .xlsx o .xls'}), 400

    try:
        from openpyxl import load_workbook
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        file.save(tmp.name)
        tmp.close()

        wb = load_workbook(tmp.name, data_only=True)
        os.unlink(tmp.name)

        conn = get_db()
        creados, actualizados = [], []

        for sheet_name in wb.sheetnames:
            if 'NUEVO' in sheet_name.upper():
                continue

            ws = wb[sheet_name]
            semi_name = sheet_name.strip()
            ingredientes = []
            rendimiento = None
            unidad_rend = 'g'

            for row in ws.iter_rows(min_row=4, values_only=True):
                if not any(row): continue
                ing_name = row[0]
                if ing_name is None: continue
                ing_str = str(ing_name).strip()

                if 'RENDIMIENTO' in ing_str.upper():
                    try:
                        rendimiento = float(row[1]) if row[1] else 0
                        unidad_rend = str(row[2]).strip() if row[2] else 'g'
                    except: pass
                    break

                if ing_str.upper() in ('INGREDIENTE','') or ing_str.startswith('🍣'):
                    continue

                try:
                    cant = float(row[1]) if row[1] is not None else 0
                    unid = str(row[2]).strip() if row[2] else 'g'
                except:
                    cant, unid = 0, 'g'

                if cant > 0 and ing_str:
                    ingredientes.append({'nombre': ing_str, 'cantidad': cant, 'unidad': unid})

            if not ingredientes and rendimiento is None:
                continue

            KEY_MAP = {
                'tartar': 'tartar', 'grill': 'grill', 'langostino': 'langos',
                'guacamole': 'guac', 'batata': 'batata', 'kanikama': 'kanikama',
                'spicy': 'spicy', 'okinawa': 'okinawa', 'salmon crispy': 'salmonCrispy',
            }
            insumo_key = semi_name.lower().replace(' ','_')
            for k, v in KEY_MAP.items():
                if k in semi_name.lower():
                    insumo_key = v
                    break

            receta_json = json.dumps(ingredientes)

            existing = conn.execute('SELECT id FROM semielaborados WHERE LOWER(name)=LOWER(?)', (semi_name,)).fetchone()
            rolls_json = json.dumps([])
            if existing:
                ex_full = conn.execute('SELECT rolls FROM semielaborados WHERE id=?', (existing['id'],)).fetchone()
                rolls_json = ex_full['rolls'] if ex_full else '[]'
                conn.execute('''UPDATE semielaborados SET insumo_key=?, unit=?, rolls=?,
                                receta=?, rendimiento_cantidad=?, rendimiento_unidad=? WHERE id=?''',
                            (insumo_key, unidad_rend, rolls_json, receta_json,
                             rendimiento or 0, unidad_rend, existing['id']))
                actualizados.append(semi_name)
            else:
                conn.execute('''INSERT INTO semielaborados
                                (name, insumo_key, unit, rolls, receta, rendimiento_cantidad, rendimiento_unidad)
                                VALUES (?,?,?,?,?,?,?)''',
                            (semi_name, insumo_key, unidad_rend, rolls_json, receta_json,
                             rendimiento or 0, unidad_rend))
                creados.append(semi_name)

        conn.commit()
        conn.close()
        return jsonify({'ok': True, 'creados': creados, 'actualizados': actualizados,
                        'total': len(creados) + len(actualizados)})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ── Main ──────────────────────────────────────────────────────────────────
init_db()

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
