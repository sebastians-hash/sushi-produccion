"""Genera el Excel de 'Insumos totales' (directos + lo que se gasta preparando
semielaborados, ya bajado a insumos crudos) con la merma de cada uno aplicada."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="5F5E5A", end_color="5F5E5A", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Arial", bold=True, size=14)
SUB_FONT = Font(name="Arial", size=10, color="666666")
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("Insumo", 32), ("Unidad", 10), ("Cantidad neta", 15),
    ("Eficiencia del insumo", 18), ("Cantidad bruta", 15),
    ("Categoría", 20), ("Proveedor", 24),
]


def build_insumos_xlsx(rows: list, date_str: str, global_pct, out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Insumos totales"

    ws.merge_cells("A1:G1")
    ws["A1"] = "INSUMOS TOTALES"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:G2")
    ws["A2"] = f"{date_str}  |  {global_pct}% de la venta"
    ws["A2"].font = SUB_FONT

    header_row = 4
    for i, (title, width) in enumerate(COLUMNS, start=1):
        col = get_column_letter(i)
        cell = ws[f"{col}{header_row}"]
        cell.value = title
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[col].width = width
    ws.row_dimensions[header_row].height = 28

    r = header_row + 1
    for row in rows:
        ws.cell(row=r, column=1, value=row["insumo"]).font = BODY_FONT
        ws.cell(row=r, column=2, value=row.get("unidad", "")).font = BODY_FONT
        c_neta = ws.cell(row=r, column=3, value=row["cantidad_neta"])
        c_neta.font = BODY_FONT
        c_neta.number_format = "#,##0.00"
        c_efic = ws.cell(row=r, column=4, value=(row.get("eficiencia", 100) or 100) / 100)
        c_efic.font = BODY_FONT
        c_efic.number_format = "0%"
        c_bruta = ws.cell(row=r, column=5, value=row["cantidad_bruta"])
        c_bruta.font = BODY_FONT
        c_bruta.number_format = "#,##0.00"
        ws.cell(row=r, column=6, value=row.get("categoria", "")).font = BODY_FONT
        ws.cell(row=r, column=7, value=row.get("proveedor", "")).font = BODY_FONT
        for col in range(1, 8):
            ws.cell(row=r, column=col).border = BORDER
        if r % 2 == 0:
            for col in range(1, 8):
                ws.cell(row=r, column=col).fill = PatternFill(start_color="F5F5F3", end_color="F5F5F3", fill_type="solid")
        r += 1

    ws.freeze_panes = f"A{header_row + 1}"
    wb.save(out_path)
    print(f"Excel generado: {out_path}")
